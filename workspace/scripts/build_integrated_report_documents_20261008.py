"""Create the integrated Word/PDF in new paths using the proven report layout."""
from pathlib import Path
import json
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.shared import Pt
import build_full_history_report as legacy_layout
from build_full_history_report import build, inline
from build_report_docx import set_run_font

ROOT=Path(__file__).resolve().parents[1]
NAME='LLM_RL火灾疏散研究全程整合报告_2026-10-08'
ASSETS=ROOT/'artifacts/integrated_report_assets_20261008'


def independent_numbering(doc, start):
    """Preserve each Markdown list start through Word's numbering normalization."""
    numbering=doc.part.numbering_part.element
    abstract_ids=[int(n.get(qn('w:abstractNumId'))) for n in numbering.findall(qn('w:abstractNum'))]
    num_ids=[int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
    abstract_id=max(abstract_ids,default=-1)+1
    num_id=max(num_ids,default=0)+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(abstract_id))
    for tag,value in (('nsid',f'{0x61080000+abstract_id:08X}'),('multiLevelType','singleLevel'),('tmpl',f'{0x62080000+abstract_id:08X}')):
        node=OxmlElement('w:'+tag);node.set(qn('w:val'),value);abstract.append(node)
    level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
    for tag,value in (('start',str(start)),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')):
        node=OxmlElement('w:'+tag);node.set(qn('w:val'),value);level.append(node)
    properties=OxmlElement('w:pPr')
    tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'720');tabs.append(tab);properties.append(tabs)
    indent=OxmlElement('w:ind');indent.set(qn('w:left'),'720');indent.set(qn('w:hanging'),'360');properties.append(indent)
    level.append(properties);abstract.append(level)
    instances=numbering.findall(qn('w:num'))
    if instances: instances[0].addprevious(abstract)
    else: numbering.append(abstract)
    instance=OxmlElement('w:num');instance.set(qn('w:numId'),str(num_id))
    reference=OxmlElement('w:abstractNumId');reference.set(qn('w:val'),str(abstract_id));instance.append(reference)
    override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0')
    restart=OxmlElement('w:startOverride');restart.set(qn('w:val'),str(start));override.append(restart);instance.append(override)
    numbering.append(instance)
    return num_id


def create_word(target=None):
    target=target or ROOT/'artifacts'/(NAME+'.docx')
    assert not target.exists()
    original_numbering=legacy_layout.ordered_numbering
    try:
        legacy_layout.ordered_numbering=independent_numbering
        metadata=build(ROOT/'reports'/(NAME+'.md'),target)
    finally:
        legacy_layout.ordered_numbering=original_numbering
    doc=Document(target)
    replaces={
        '全流程实验与方法复盘报告':'全程整合研究报告',
        '到动态门控、公平对照、护栏审计与计划蒸馏':'到收益预测、候选审计与再次覆盖门控验证',
        '截至2026年10月2日的阶段性证据':'截至2026年10月8日的完整阶段记录',
    }
    for p in doc.paragraphs[:25]:
        if p.text in replaces:
            for r in p.runs:
                if r.text in replaces: r.text=replaces[r.text]
    for section in doc.sections:
        for p in section.header.paragraphs:
            for r in p.runs:
                if '全流程实验复盘｜2026-10-02' in r.text:
                    r.text=r.text.replace('全流程实验复盘｜2026-10-02','全程整合报告｜2026-10-08')
    # Replace the legacy flat navigation with a real Word table-of-contents field.
    paragraphs=doc.paragraphs
    nav=next(i for i,p in enumerate(paragraphs) if p.text=='报告导航')
    body=next(i for i in range(nav+1,len(paragraphs)) if paragraphs[i].style.name.startswith('Heading'))
    for p in paragraphs[nav+1:body]: p._element.getparent().remove(p._element)
    navigation=paragraphs[nav]
    navigation.style=doc.styles['Title']
    set_run_font(navigation.runs[0],chinese='微软雅黑',size=16,bold=True)
    toc=paragraphs[body].insert_paragraph_before()
    toc.paragraph_format.first_line_indent=Pt(0)
    r=toc.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    r=toc.add_run();instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-1" \\h \\z \\u ';r._r.append(instruction)
    r=toc.add_run();separator=OxmlElement('w:fldChar');separator.set(qn('w:fldCharType'),'separate');r._r.append(separator)
    set_run_font(toc.add_run('目录将在PDF导出前由Word更新。'),size=10.5)
    r=toc.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
    paragraphs[body].insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)
    header=(ROOT/'reports'/(NAME+'.md')).read_text(encoding='utf-8').split('## 阅读说明与当前结论')[0]
    for line in header.splitlines():
        if line.startswith('**整合日期') or line.startswith('用途：'):
            intro=paragraphs[body].insert_paragraph_before();inline(intro,line)
    # Chapter titles begin fresh pages; small subdivisions remain with their text.
    for p in doc.paragraphs:
        if p.style.name=='Heading 1' and p.text not in ('报告导航','阅读说明与当前结论'):
            p.paragraph_format.page_break_before=True
    doc.core_properties.title='LLM与强化学习火灾疏散研究全程整合报告（2026-10-08）'
    doc.core_properties.subject='完整研究经过、逐轮方法与正负结果、机制诊断、当前状态及证据索引'
    doc.save(target)
    (ASSETS/'word_build_manifest.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(metadata,ensure_ascii=False),flush=True)
    return target


def export_pdf(target):
    import win32com.client
    pdf=target.with_suffix('.pdf');assert not pdf.exists()
    app=win32com.client.DispatchEx('Word.Application')
    document=None
    try:
        app.Visible=False;app.DisplayAlerts=0;app.AutomationSecurity=3
        document=app.Documents.Open(str(target),ConfirmConversions=False,ReadOnly=False,AddToRecentFiles=False)
        document.Fields.Update()
        for i in range(1,document.TablesOfContents.Count+1): document.TablesOfContents(i).Update()
        document.Repaginate()
        for i in range(1,document.TablesOfContents.Count+1): document.TablesOfContents(i).UpdatePageNumbers()
        document.Save()
        document.ExportAsFixedFormat(str(pdf),17,OpenAfterExport=False,OptimizeFor=0,CreateBookmarks=1)
        print(json.dumps(dict(pdf=str(pdf),toc_count=document.TablesOfContents.Count,pages=document.ComputeStatistics(2)),ensure_ascii=False),flush=True)
    finally:
        if document is not None: document.Close(False)
        app.Quit()


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');export_pdf(create_word())

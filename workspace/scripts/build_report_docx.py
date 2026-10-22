from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "LLM_RL火灾疏散研究与实验实施报告.md"
OUTPUT = ROOT / "artifacts" / "LLM_RL火灾疏散研究与实验实施报告.docx"
ASSETS = ROOT / "artifacts" / "report_assets"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, chinese="宋体", latin="Times New Roman", size=11, bold=None, color="000000") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color="666666")


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Pt(22)

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("基于真实行为校准的LLM疏散者与\n强化学习动态引导研究报告")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("研究方案与第一阶段代码实践")
    set_run_font(r, chinese="微软雅黑", size=15, bold=False, color="333333")
    doc.add_paragraph()
    doc.add_paragraph()
    for label, value in (
        ("研究方向", "LLM行为建模  强化学习控制  火灾疏散"),
        ("报告用途", "开题论证  导师讨论  实验实施"),
        ("代码版本", "0.6.0 自适应稳健训练原型"),
        ("日期", "2026年9月28日"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}  {value}")
        set_run_font(r, chinese="微软雅黑", size=11.5, color="333333")
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("报告结构", level=1)
    items = [
        "研究背景与问题界定",
        "研究目标与科学问题",
        "总体技术路线",
        "真实数据与行为模型实验",
        "火灾与疏散仿真设计",
        "RL动态引导设计",
        "实验矩阵与评价指标",
        "稳健性与消融实验",
        "防止LLM自证循环的质量控制",
        "第一阶段代码实践",
        "下一阶段代码路线",
        "预期论文贡献、计划与风险",
    ]
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        r = p.add_run(f"{i}.  {item}")
        set_run_font(r, size=11)
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    header_tr_pr.append(tbl_header)
    for j, value in enumerate(rows[0]):
        cell = table.rows[0].cells[j]
        cell.text = value
        set_cell_shading(cell, "365F7D")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, chinese="微软雅黑", size=9.5, bold=True, color="FFFFFF")
    for i, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        row_pr = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_pr.append(cant_split)
        for j, value in enumerate(row):
            cell = cells[j]
            cell.text = value
            if i % 2:
                set_cell_shading(cell, "F3F7FA")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 or re.fullmatch(r"[\d.%-]+", value) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, chinese="等线", latin="Consolas", size=8.5, color="222222")


def add_figure(doc: Document, filename: str, caption: str, width: float = 6.8) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    r = cap.add_run(caption)
    set_run_font(r, size=9, color="555555")


def add_inline_markdown(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, size=11, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, chinese="等线", latin="Consolas", size=9.5, color="333333")
        else:
            r = paragraph.add_run(part)
            set_run_font(r, size=11)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)
    add_cover(doc)
    add_contents(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 摘要"))
    i = start
    inserted_architecture = False
    inserted_results = False
    inserted_ppo_results = False
    inserted_multiseed_results = False
    inserted_cross_results = False
    inserted_recurrent_results = False
    inserted_robust_results = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            doc.add_heading(title, level=1)
            i += 1
            continue
        if line.startswith("### "):
            title = line[4:].strip()
            doc.add_heading(title, level=2)
            if title == "3.1 模块分工" and not inserted_architecture:
                add_figure(doc, "architecture.png", "图1 真实行为校准与动态引导的闭环实验架构")
                inserted_architecture = True
            if title == "10.4 原型结果摘要" and not inserted_results:
                add_figure(doc, "prototype_results.png", "图2 第一阶段合成原型结果 失败比例经放大显示")
                inserted_results = True
            if title == "10.7 PPO实现与第二阶段结果" and not inserted_ppo_results:
                add_figure(doc, "ppo_results.png", "图3 合成校准Logit人群中的PPO结果 失败比例经放大显示")
                inserted_ppo_results = True
            if title == "10.8 五训练种子与分层Bootstrap" and not inserted_multiseed_results:
                add_figure(doc, "ppo_multiseed_results.png", "图4 五个PPO训练种子在合成校准Logit人群上的失能率")
                inserted_multiseed_results = True
            if title == "10.9 状态历史与行为交叉训练实验" and not inserted_cross_results:
                add_figure(doc, "behavior_cross_heatmap.png", "图5 PPO训练行为与测试行为的失能率交叉矩阵")
                inserted_cross_results = True
            if title == "10.10 历史堆叠与RecurrentPPO对照" and not inserted_recurrent_results:
                add_figure(doc, "recurrent_comparison.png", "图6 4步历史PPO与RecurrentPPO失能率对照")
                inserted_recurrent_results = True
            if title == "10.11 最坏情况自适应行为集成" and not inserted_robust_results:
                add_figure(doc, "robust_curriculum.png", "图7 三个训练种子的行为采样权重阶段均值")
                add_figure(doc, "robust_methods.png", "图8 三种训练方法的平均与最坏行为失能率")
                inserted_robust_results = True
            i += 1
            continue
        if line.startswith("```"):
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code(doc, code_lines)
            i += 1
            continue
        if line == "\\[":
            equation_lines = []
            i += 1
            while i < len(lines) and lines[i] != "\\]":
                equation_lines.append(lines[i])
                i += 1
            latex = "".join(equation_lines)
            filename = "equation_state.png" if "s_t=" in latex else "equation_reward.png"
            add_figure(doc, filename, "", width=5.8)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_rows: list[list[str]] = []
            table_rows.append([c.strip() for c in line.strip("|").split("|")])
            i += 2  # skip separator
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, table_rows)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline_markdown(p, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            match = re.match(r"^(\d+)\.\s+(.*)$", line)
            assert match is not None
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            add_inline_markdown(p, f"{match.group(1)}.  {match.group(2)}")
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or nxt.startswith(("#", "- ", "|", "```", "\\[")) or re.match(r"^\d+\. ", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline_markdown(p, " ".join(paragraph_lines))

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = header.add_run("LLM与强化学习火灾疏散研究")
        set_run_font(r, chinese="微软雅黑", size=8.5, color="777777")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())


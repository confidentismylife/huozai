"""Typeset the Chinese v3 manuscript using the user's reference-paper layout.

Run with Codex bundled Python. Reads only the frozen analysis outputs. Does not
call model APIs, train, or simulate. Native Word math is used for three equations.
"""
from pathlib import Path
import hashlib
import json
import re
import shutil

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/llm_ppo_paper_v3_20261021'
OLD=ROOT/'artifacts/llm_ppo_paper_v2_20261021'
MD=OUT/'固定LLM建议与PPO疏散控制_期刊式改写稿_v3.md'
DOCX=MD.with_suffix('.docx')
REF=Path('D:/UserFolders/Desktop/火灾行人轨迹疏散论文呢/1-s2.0-S2352710225033698-main的全文翻译.pdf')
WIDTH=468

def font(run,size=10.5,bold=False,east='宋体'):
    run.font.name='Times New Roman'
    run.font.size=Pt(size); run.font.bold=bold; run.font.italic=False
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)
    run.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)

def text(p,value,size=10.5,bold=False):
    # Strip markdown emphasis without changing the scientific text.
    value=value.replace('**','').replace('`','')
    run=p.add_run(value); font(run,size,bold); return run

def border(el,kind='bottom',size='4',color='000000'):
    node=OxmlElement('w:'+kind)
    node.set(qn('w:val'),'single'); node.set(qn('w:sz'),size); node.set(qn('w:color'),color)
    el.append(node)

def format_paragraph(p,indent=True):
    p.paragraph_format.line_spacing=Pt(14.2)
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.first_line_indent=Pt(21 if indent else 0)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control=True

def table(doc,rows,widths=None,front=False):
    t=doc.add_table(rows=1,cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths is None:
        n=len(rows[0]); widths=([155]+[(WIDTH-155)/(n-1)]*(n-1)) if n>1 else [WIDTH]
    for col,w in zip(t.columns,widths): col.width=Pt(w)
    props=t._tbl.tblPr
    edges=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'nil'); edges.append(e)
    if not front:
        border(edges,'top','8'); border(edges,'bottom','8')
    props.append(edges)
    for i,values in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        row=t.rows[i]
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if i==0 and not front: trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Pt(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            tcpr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for direction,amt in [('top',65),('bottom',65),('left',65),('right',65)]:
                node=OxmlElement('w:'+direction); node.set(qn('w:w'),str(amt)); node.set(qn('w:type'),'dxa'); mar.append(node)
            tcpr.append(mar)
            if i==0 and not front:
                bd=OxmlElement('w:tcBorders'); border(bd,'bottom','4'); tcpr.append(bd)
            if front:
                for k,part in enumerate(value.split('\n')):
                    if not part: continue
                    p=cell.paragraphs[0] if k==0 else cell.add_paragraph()
                    format_paragraph(p,False)
                    p.paragraph_format.line_spacing=Pt(13)
                    p.paragraph_format.space_after=Pt(6 if k==0 else 2)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or k==0 else WD_ALIGN_PARAGRAPH.JUSTIFY
                    text(p,part,9.5,k==0)
                continue
            p=cell.paragraphs[0]; format_paragraph(p,False)
            # Keep each short scientific table together and never leave only a
            # header at the foot of a page. Longer tables retain repeated heads.
            p.paragraph_format.keep_with_next=(i==0 or (len(rows)<=6 and i<len(rows)-1))
            p.paragraph_format.line_spacing=Pt(12.5 if not front else 13)
            p.paragraph_format.space_after=Pt(2)
            if front: p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            else: p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(value)>22 else WD_ALIGN_PARAGRAPH.CENTER
            text(p,value,9.2 if not front else 9.5,i==0 and not front)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def math_run(value):
    r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=value; r.append(t); return r

def sub(base,index):
    node=OxmlElement('m:sSub'); b=OxmlElement('m:e'); b.append(math_run(base)); s=OxmlElement('m:sub'); s.append(math_run(index)); node.extend([b,s]); return node

def equation(doc,number):
    p=doc.add_paragraph(); format_paragraph(p,False)
    p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(460),WD_TAB_ALIGNMENT.RIGHT)
    om=OxmlElement('m:oMath')
    if number==1:
        om.append(math_run('𝒜 = {N, S} × {0.25, 0.50, 0.75, 1.00} ∪ {wait}'))
    elif number==2:
        om.extend([sub('D','i'),math_run('(t + 1) = '),sub('D','i'),math_run('(t) + '),sub('h','zᵢ(t)'),math_run('(t)')])
    else:
        om.extend([sub('r','t'),math_run(' = ΔEₜ/P − 20ΔIₜ/P − 0.2ΔDₜ/P − 0.1Cₜ − 0.01Sₜ')])
    p._p.append(om); text(p,'\t('+str(number)+')',10)

def heading(doc,value,level):
    p=doc.add_paragraph(style='Heading '+str(level)); text(p,value,12 if level==1 else 10.8,True)
    p.paragraph_format.space_before=Pt(10 if level==1 else 7); p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True; p.paragraph_format.first_line_indent=Pt(0)
    return p

def build():
    OUT.mkdir(parents=True,exist_ok=True); (OUT/'figures').mkdir(exist_ok=True)
    for file in (OLD/'figures').glob('*.png'): shutil.copy2(file,OUT/'figures'/file.name)
    source_hash=hashlib.sha256(MD.read_bytes()).hexdigest()
    lines=MD.read_text(encoding='utf-8').splitlines()
    doc=Document(); section=doc.sections[0]
    section.page_width=Pt(544); section.page_height=Pt(742)
    section.left_margin=Pt(38); section.right_margin=Pt(38)
    section.top_margin=Pt(40); section.bottom_margin=Pt(40)
    section.header_distance=Pt(19); section.footer_distance=Pt(19)
    normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name]; style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    # Word's built-in Title can carry an inherited blue bottom border.
    # Remove that decoration so the heading follows the supplied journal sample.
    for node in list(doc.styles.element.xpath('.//w:pBdr')):
        node.getparent().remove(node)
    settings=doc.settings.element; update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
    header=section.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    text(header,'固定LLM建议与PPO火灾疏散控制  |  中文写作稿',7.5)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    title=doc.add_paragraph(style='Title'); text(title,lines[0][2:],18,True); title.paragraph_format.space_after=Pt(7)
    title.paragraph_format.keep_with_next=True
    subtitle=doc.add_paragraph(style='Subtitle'); text(subtitle,lines[2][3:],11.5,False); subtitle.paragraph_format.space_after=Pt(7); subtitle.paragraph_format.keep_with_next=True
    p=doc.add_paragraph(); text(p,'中文论文写作稿    2026年10月21日',8.5); p.paragraph_format.space_after=Pt(10)
    a=lines.index('**摘要**'); abstract=lines[a+2]
    keyword=next(s for s in lines if s.startswith('**关键词**'))
    table(doc,[['关键词\n\n'+keyword.replace('**关键词** ','').replace('；','\n'),'摘  要\n\n'+abstract]],widths=[108,360],front=True)
    start=next(i for i,line in enumerate(lines) if line.startswith('## 1 '))
    i=start; references=False
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line=='$$':
            block=[]; i+=1
            while lines[i].strip()!='$$': block.append(lines[i]); i+=1
            number=int(re.search(r'\\tag\{(\d+)\}',' '.join(block)).group(1)); equation(doc,number); i+=1; continue
        if line.startswith('### '): heading(doc,line[4:],2); i+=1; continue
        if line.startswith('## '):
            value=line[3:]; heading(doc,value,1); references=value=='参考文献'; i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in cells): rows.append(cells)
                i+=1
            widths={1:[120,114,120,114],2:[83,225,160],3:[100,38,68,72,105,85],4:[94,94,88,96,96],5:[95,114,62,121,76]}
            number=len(doc.tables) # includes frontmatter
            table(doc,rows,widths=widths.get(number))
            continue
        if line.startswith('!['):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line); label,path=match.groups()
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
            image=OUT/path
            with Image.open(image) as im: ratio=im.height/im.width
            width=min(WIDTH,235/ratio)
            p.add_run().add_picture(str(image),width=Pt(width))
            caption=doc.add_paragraph(); caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            text(caption,label,9); caption.paragraph_format.space_after=Pt(6)
            i+=1; continue
        p=doc.add_paragraph(); format_paragraph(p,not references)
        if re.match(r'^表\d+ ',line):
            p.paragraph_format.keep_with_next=True; p.paragraph_format.first_line_indent=Pt(0); text(p,line,9.5,True)
        elif references:
            p.paragraph_format.left_indent=Pt(18); p.paragraph_format.first_line_indent=Pt(-18)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing=Pt(11.5); text(p,line,8.6)
        else: text(p,line)
        i+=1
    doc.core_properties.title=lines[0][2:]
    doc.core_properties.subject='固定LLM建议与PPO群体疏散控制的对照和机制分析'
    doc.core_properties.author=''
    doc.core_properties.comments='参考用户提供论文的章节层次与版式；非期刊已录用稿。'
    doc.save(DOCX)
    evidence={
      'status':'assistant_checked_working_draft_requires_author_review',
      'manuscript_sha256':source_hash,'reference_pdf_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),
      'layout_reference':{'page_points':[544,742],'body':'single column','frontmatter':'keywords alongside abstract','headings':'seven numbered sections','tables':'three horizontal rules','publisher_branding':'not reproduced'},
      'source_documents':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [
        OLD/'固定LLM建议对PPO疏散控制的影响_中文论文_v2.md',OLD/'data/论文主表_修订.csv',
        OLD/'data/endpoint_attainment_summary.csv',OLD/'data/conditional_paired_times.csv',
        OLD/'data/prediction_without_identity_metrics.csv',OLD/'data/primary_absolute_outcomes.csv']},
      'figure_sources':{p.name:str((OLD/'figures'/p.name).relative_to(ROOT)) for p in (OUT/'figures').glob('*.png')},
      'new_ppo_training':0,'new_rollouts':0,'new_llm_requests':0,
      'human_verification':'pending; no automated result is labelled human-verified',
      'citation_support':{
        '1':{'source':'user-supplied full reference PDF and publisher-deposited Crossref metadata','doi':'10.1016/j.jobe.2025.115132','scope':'SafeStep general method and bibliographic identity'},
        '2':{'source':'https://arxiv.org/abs/1707.06347','scope':'PPO objective and author metadata'},
        '3':{'source':'https://aclanthology.org/2025.acl-long.1438/','scope':'FLARE abstract and bibliography'},
        '4':{'source':'https://api.crossref.org/works/10.1016/j.ssci.2025.106935','scope':'title-level research topic and bibliographic metadata; full text not reviewed this turn'},
        '5':{'source':'https://api.crossref.org/works/10.1016/j.ress.2025.112056','scope':'title-level research topic and bibliographic metadata; full text not reviewed this turn'},
        '6':{'source':'https://arxiv.org/abs/2108.13264','scope':'uncertainty of limited-run RL evaluation and author metadata'}}}
    (OUT/'evidence_and_layout_record.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(str(DOCX).encode('ascii','backslashreplace').decode())

if __name__=='__main__': build()

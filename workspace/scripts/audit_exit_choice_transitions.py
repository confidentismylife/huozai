"""Read-only source audit and cached diagnostics. No training or inference."""
from collections import Counter, defaultdict
from datetime import date, datetime, time, timezone
import csv
import hashlib
import json
from pathlib import Path
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PREV = ROOT / "outputs/exit_choice_initial_benchmark_20261012"
OUT = ROOT / "outputs/exit_choice_transition_audit_20261012"
XLSX = Path("C:/Users/摩天轮大盗/Downloads/Exit choice.xlsx")
PPTX = Path("C:/Users/摩天轮大盗/Downloads/Exit choice ppt.pptx")


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def read(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))


def save(name, value):
    (OUT/name).write_text(json.dumps(value, ensure_ascii=False, indent=2,
                                   allow_nan=False, default=str), encoding="utf-8")


def csvsave(name, rows):
    with (OUT/name).open("w", encoding="utf-8-sig", newline="") as f:
        if rows:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            for r in rows:
                w.writerow({k: json.dumps(v, ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in r.items()})


def cellvalues(ws, row):
    return {c.coordinate: c.value for c in ws[row] if c.value is not None}


def is_clock_value(v):
    return isinstance(v, (datetime,time)) or isinstance(v,str) and bool(re.fullmatch(r"\s*\d{1,2}:\d{2}(?::\d{2})?\s*",v))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    # Live Ollama logs may change independently. Protect research files, all raw
    # answers, both attachments and previous reports, without accessing held-out scenes.
    protected = [XLSX, PPTX] + [p for p in PREV.rglob("*") if p.is_file() and p.suffix != ".log"]
    protected += [ROOT/"reports/Exit_choice三出口数据清洗与首轮实验_2026-10-12.md",
                  ROOT/"reports/Exit_choice三出口字段与清洗规则_2026-10-12.md",
                  ROOT/"outputs/local_llm_selection.json", ROOT/"outputs/third_round_development_status.json"]
    before = [{"path":str(p), "sha256":sha(p)} for p in protected]
    save("preservation_before.json", before)
    wb = openpyxl.load_workbook(XLSX, read_only=False, data_only=False)
    raw, sequences = read(PREV/"raw_level2_records.json"), read(PREV/"exploratory_sequence_cases.json")
    pairs, review = [], []
    for case in sorted(sequences, key=lambda r:(r["date"],r["row"])):
        records = sorted([r for r in raw if r["person_key"]==case["person_key"]],key=lambda r:r["row"])
        assert len(records)==2
        a,b = records
        assert a["initial_label"]==1 and a["final_label"]==2 and b["initial_label"]==2
        ws = wb[a["sheet"]]
        cells_a, cells_b = cellvalues(ws,a["row"]), cellvalues(ws,b["row"])
        block = next(r for r in read(PREV/"block_mapping.json") if r["sheet"]==a["sheet"] and r["header_row"]==a["header_row"])
        end_col = block["choice"]+1
        extra_a = {k:v for k,v in cells_a.items() if openpyxl.utils.cell.coordinate_from_string(k)[0] and openpyxl.utils.column_index_from_string(openpyxl.utils.cell.coordinate_from_string(k)[0])>end_col}
        extra_b = {k:v for k,v in cells_b.items() if openpyxl.utils.column_index_from_string(openpyxl.utils.cell.coordinate_from_string(k)[0])>end_col}
        notes = [str(v) for v in extra_a.values()]+[str(a["raw_choice"])]
        # These notes describe the observed outcome. They are never a new input cue.
        attributed = any("staff" in v.lower() and "due" in v.lower() for v in notes)
        fa = {field:[e[field] for e in a["exits"]] for field in ("dist","people","staff","fam")}
        fb = {field:[e[field] for e in b["exits"]] for field in ("dist","people","staff","fam")}
        clock_a = {k:str(v) for k,v in cells_a.items() if is_clock_value(v)}
        clock_b = {k:str(v) for k,v in cells_b.items() if is_clock_value(v)}
        issues = ["no_person_aligned_pre_choice_cue_and_post_choice_timestamps",
                  "no_matched_repeated_observation_for_non_changers",
                  "Staff_and_People_unchanged_between_records"]
        if attributed and fa["staff"]==[0,0,0] and fb["staff"]==[0,0,0]:
            issues.append("staff_attribution_with_all_zero_Staff_covariates")
        group5 = bool(re.search(r"group\s+of\s+5",a["description_for_matching_only"],re.I))
        if group5:
            issues.append("explicit_group_of_five_shared_observations")
        pair = {
            "person_key":a["person_key"], "date":a["date"], "person_id":a["person_id"],
            "sheet":a["sheet"], "rows":[a["row"],b["row"]],
            "choice_cells":[a["choice_cell"],b["choice_cell"]],
            "choices_raw":[a["raw_choice"],b["raw_choice"]],
            "initial_exit":1, "later_exit":2,
            "earliest_case_id":a["record_id"], "later_case_id":b["record_id"],
            "recorded_before_attributes":fa, "recorded_later_attributes":fb,
            "distance_delta":[y-x for x,y in zip(fa["dist"],fb["dist"])],
            "unchanged_fields":[f for f in fa if fa[f]==fb[f]],
            "extra_cells_earliest":extra_a, "extra_cells_later":extra_b,
            "source_note_attributes_change_to_staff":attributed,
            "person_clock_cells_earliest":clock_a, "person_clock_cells_later":clock_b,
            "cue_timestamp":None, "initial_choice_timestamp":None, "later_choice_timestamp":None,
            "source_video_identifier":None, "reaction_latency_seconds":None,
            "explicit_group_of_five":group5,
            "group_key": "2019-03-06:meeting_room_group5" if group5 else None,
            "endpoint_sequence_supported":True,
            "fully_time_aligned_predictive_pair":False,
            "static_attribute_alignment_before_first_choice_verified":False,
            "issues":issues,
            "ppt_slide":a["ppt_slide"],
        }
        pairs.append(pair)
        review.append({"person_key":a["person_key"],"sheet":a["sheet"],"source_rows":pair["rows"],
                       "initial_exit":1,"later_exit":2,"recorded_staff_before":fa["staff"],
                       "recorded_staff_later":fb["staff"],"source_note_staff_attribution":attributed,
                       "source_video_identifier":"","initial_choice_time":"","staff_cue_time":"",
                       "post_cue_observation_time":"","later_choice_time":"","Staff_encoding_verified":"",
                       "non_changer_comparison_identifier":"","group_key":pair["group_key"],
                       "review_status":"unverified_timing", "source_issues":issues})
    save("transition_pairs.json",pairs)
    csvsave("transition_review.csv",review)
    csvsave("transition_summary.csv",[{"person_key":p["person_key"],"date":p["date"],"rows":p["rows"],
               "initial_exit":p["initial_exit"],"later_exit":p["later_exit"],
               "distance_before":p["recorded_before_attributes"]["dist"],
               "distance_later":p["recorded_later_attributes"]["dist"],
               "people_unchanged":"people" in p["unchanged_fields"],
               "staff_unchanged":"staff" in p["unchanged_fields"],
               "staff_change_attribution":p["source_note_attributes_change_to_staff"],
               "all_staff_zero":p["recorded_before_attributes"]["staff"]==[0,0,0],
               "explicit_group_of_five":p["explicit_group_of_five"],
               "time_aligned_predictive_pair":p["fully_time_aligned_predictive_pair"]} for p in pairs])
    # Audit raw codes across all sheets, using actual Staff header columns.
    code_counts, unusual_codes, time_clues = Counter(), [], []
    for ws in wb:
        staff_cols = []
        for row in ws:
            vals = [c.value for c in row]
            if any(isinstance(v,str) and v.strip().lower()=="dist" for v in vals):
                staff_cols = [c.column for c in row if isinstance(c.value,str) and c.value.strip().lower()=="staff"]
            identity = vals[0]
            if identity is not None and re.fullmatch(r"\d+(?:\.0)?[a-z]?",str(identity),re.I):
                for col in staff_cols:
                    cell = ws.cell(row[0].row,col)
                    code_counts[str(cell.value)]+=1
                    if cell.value not in (0,1,None):
                        unusual_codes.append({"sheet":ws.title,"row":cell.row,"person_id":str(identity),
                                              "cell":cell.coordinate,"raw_code":cell.value})
            elif any(is_clock_value(v) or isinstance(v,str) and any(k in v.lower() for k in ["hesitation","evacuation time","directed by staff"]) for v in vals):
                time_clues.append({"sheet":ws.title,"row":row[0].row,
                                   "relevant_cells":{c.coordinate:str(c.value) for c in row if is_clock_value(c.value) or isinstance(c.value,str) and any(k in c.value.lower() for k in ["hesitation","evacuation time","directed by staff"])},
                                   "usable_as_GF_individual_reaction_time":False})
    roles = []
    for sheet,ident in [("06-03-19 EC","69"),("11-09-19 EC","125"),("11-09-19 EC","126"),("06-03-19 EC","70")]:
        r = next(r for r in raw if r["sheet"]==sheet and r["person_id"]==ident)
        roles.append({"person_key":r["person_key"],"sheet":sheet,"row":r["row"],
                      "description":r["description_for_matching_only"],
                      "staff_codes":[e["staff"] for e in r["exits"]],
                      "identity_staff_annotated":r["staff_role"],
                      "explicit_preparation_to_alert_note":ident=="126",
                      "evacuation_management_role_time_verified":False,
                      "prior_exclusion":r["staff_role"],
                      "interpretation": "staff_identity_or_group_annotated_role_timing_unresolved" if r["staff_role"] else "adjacent_group_description_not_sufficient_to_assign_staff_identity"})
    wb.close()
    save("staff_role_audit.json",roles)
    csvsave("staff_role_audit.csv",roles)
    save("staff_code_audit.json",{"counts":dict(code_counts),"unusual_codes":unusual_codes,
                                 "interpretation_of_code_2_confirmed":False,
                                 "no_automatic_2_to_binary_mapping":True})
    save("time_clues.json",time_clues)
    ns={"a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(PPTX) as z:
        slide_evidence=[]
        for i in (1,5,10):
            xml=ET.fromstring(z.read(f"ppt/slides/slide{i}.xml"))
            texts=["".join(c.text or "" for c in p.findall('.//a:t',ns)) for p in xml.findall('.//a:p',ns)]
            slide_evidence.append({"slide":i,"non_numeric_texts":[t for t in texts if t.strip() and not re.fullmatch(r"\d+[a-z]?",t.strip(),re.I)],
                                   "explicit_clock_texts":[t for t in texts if is_clock_value(t)]})
        media = [s for s in z.namelist() if s.startswith("ppt/media/")]
        embedded_videos=[s for s in media if Path(s).suffix.lower() in {'.mp4','.avi','.mov','.wmv','.mkv'}]
    save("ppt_annotation_evidence.json",{"slides":slide_evidence,"embedded_video_files":embedded_videos,
          "arrow_legend_is_posthoc_outcome_annotation_not_predecision_input":True,
          "rendered_slide5_for_visual_review":"source_slide_5.png",
          "existing_slide10_review_source":str(ROOT/"outputs/exit_choice_dataset_inspection_20261011/source_slide_10.png")})
    preds=read(PREV/"all_predictions.json")
    case_map={p["earliest_case_id"]:p for p in pairs}
    diagnostic=[]
    for r in preds:
        if r["cohort"]!="sequence_exploratory" or r["channel"]!="primary" or r["case_id"] not in case_map or r["method"]=="uniform":
            continue
        p=case_map[r["case_id"]]
        target=max(range(3),key=lambda j:r["probabilities"][j])+1
        diagnostic.append({"person_key":p["person_key"],"case_id":r["case_id"],"method":r["method"],
                           "input_stage":"earliest_record_only","probabilities":r["probabilities"],
                           "predicted_exit":target,"initial_exit":1,"later_exit":2,
                           "matches_initial":target==1,"matches_later":target==2,
                           "is_dynamic_prediction_validation":False})
    csvsave("cached_stage_diagnostic.csv",diagnostic)
    dsummary=[]
    for method in ("training_frequency","nearest_exit","logit","qwen3_8b"):
        rr=[r for r in diagnostic if r["method"]==method]
        assert len(rr)==15
        dsummary.append({"method":method,"n":15,"matches_initial":sum(r["matches_initial"] for r in rr),
                         "matches_later":sum(r["matches_later"] for r in rr),
                         "other_exit":sum(r["predicted_exit"]==3 for r in rr)})
    save("cached_stage_diagnostic_summary.json",dsummary)
    unique_profiles={json.dumps([p["date"],p["recorded_before_attributes"],p["recorded_later_attributes"]],sort_keys=True) for p in pairs}
    summary={"pair_count":len(pairs),"events":len({p["date"] for p in pairs}),
             "changes_1_to_2":len(pairs),"observed_paired_stayers":0,
             "staff_attributed_changes":sum(p["source_note_attributes_change_to_staff"] for p in pairs),
             "distance_changed_pairs":sum("dist" not in p["unchanged_fields"] for p in pairs),
             "E1_distance_decreased_pairs":sum(p["distance_delta"][0]<0 for p in pairs),
             "people_unchanged_pairs":sum("people" in p["unchanged_fields"] for p in pairs),
             "staff_unchanged_pairs":sum("staff" in p["unchanged_fields"] for p in pairs),
             "staff_attribution_with_all_zero_staff": [p["person_key"] for p in pairs if p["source_note_attributes_change_to_staff"] and p["recorded_before_attributes"]["staff"]==[0,0,0]],
             "explicit_group_of_five_records":sum(p["explicit_group_of_five"] for p in pairs),
             "unique_date_and_attribute_pair_profiles":len(unique_profiles),
             "fully_time_aligned_predictive_pairs":0,"paired_person_clock_values":sum(bool(p["person_clock_cells_earliest"] or p["person_clock_cells_later"]) for p in pairs),
             "embedded_video_count":len(embedded_videos),"extra_llm_calls":0,"extra_training_runs":0,
             "can_validate_changed_vs_stayed_decision":False,
             "can_estimate_reaction_latency":False,"can_estimate_causal_guidance_effect":False,
             "original_control_weights_and_frozen_scenarios_accessed":False,
             "cached_stage_diagnostic":dsummary}
    assert summary["pair_count"]==15 and summary["staff_attributed_changes"]==15
    assert summary["people_unchanged_pairs"]==15 and summary["staff_unchanged_pairs"]==15
    save("audit_summary.json",summary)
    lines=["# 改道序列时点与工作人员编码核查", "", "2026年10月12日。基于原始Excel、PPT和已保存预测，未新增模型调用或训练。", "",
           "**15条序列能确认记录中的出口1到出口2变化，但目前有0条满足本轮动态预测所需的完整时点核实条件。** 关键缺项是工作人员提醒与前后决策的时间对应，以及相同观察窗口内保持原路线者的重复观测。", "",
           "## 本轮完成与发现", "",
           "逐条配对15个编号的30行原始记录，检查CHOICE、额外备注、距离、人群计数和Staff；复查PPT第1页图例、第5和第10页地图标注，以及3个Staff描述编号。",
           "| 检查内容 | 实际结果 | 含义 |", "|---|---|---|",
           "| 选择序列 | 15条全部1转2，来自2个日期 | 只能观察这一类已发生的变化 |",
           "| 保持原路线的对应前后观测 | 0条已建立 | 不能评价误改道或学习是否该改道 |",
           "| 距离前后变化 | 15条都有变化，E1距离均减小 | 支持位置发生变化，无法单凭距离确定准确改道时刻 |",
           "| People变化 | 15条均不变 | 没有记录到可检验的人群信息变化 |",
           "| Staff变化 | 15条均不变 | 未记录出提醒信号从无到有的过程 |",
           "| 个体决策与提醒时间 | 15条均未在对应行找到 | 不能计算反应时间或判断输入先于结果 |",
           "| 同行群体 | 37至41号明确为同一5人群体，前后属性相同 | 不能把这些当作5个独立情境 |", "",
           "People和Staff不变可能是变化很快、编码概括整个阶段或其他记录方式。现有资料不足以确定原因，也不能断言原作者复制了字段。", "",
           "## 一个需要优先核对的矛盾", "",
           "06-03-19第99行，46号的CHOICE为“1 then 2”，P99备注“due to staff”；第99、100行三个出口的Staff却均为0。",
           "这说明事后归因备注与原始Staff属性的关系不能自动按二元提醒事件处理。可能涉及时点、工作人员位置、变量定义或编码问题；目前没有依据选择其中一种解释。",
           "其余14条的Staff在前后均为[0,1,0]，仍不能从中确定工作人员是在首次选择之前提醒，还是在行进中提醒。不能用“due to staff”直接补成输入Guide=1。", "",
           "## 时间和地图核查", "",
           "PPT图例说明绿色箭头表示“受到工作人员提醒因此改变选择”。它支持原标注的行为归因，没有提供逐人提醒时间。第5、第10页也没有明确个体时钟文本，PPT内未发现嵌入视频。",
           "11-09-19顶部A2、B2分别有09:52:14和07:59:48，缺少与15个前后决策一一对应的标注，不能拿来计算反应时间。",
           "工作簿其他楼层存在“Evacuation time in video”，另有31-10-19“NO HESITATION TIME FOR THIS DAY”；这些不能替代当前GF序列的首次、提醒和再次选择时点。已检查项目资料内没有找到可用于这15条核实的原始视频。", "",
           "## 工作人员身份与Staff=2", "",
           "| 日期与编号 | 来源 | 本轮判断 |", "|---|---|---|",
           "| 03-06，69 | B135：Group of 2 staff blue hi-vis；Staff=[1,1,0] | 描述涉及工作人员和两人群体，管理职责及两人编号对应未核实 |",
           "| 09-11，125 | B219：Staff left；Staff=[2,0,0] | 工作人员身份有文字依据，是否属于被排除的管理人员未核实 |",
           "| 09-11，126 | B220：Staff right；C220：Grab hi-vest to alert students；Staff=[2,0,0] | 有准备提醒学生的备注，具体时间和样本筛选角色未核实 |",
           "| 03-06，70（相邻核查） | B136：Group of 2 european lady | 不能因相邻或同样写Group of 2就自动认定是工作人员 |", "",
           "全工作簿Staff非0/1且非空的值仅发现2，位于F219、F220。2可能涉及人数或特殊编码，原论文的Guide二元定义没有解释原表2的转换，尚不能确定。保持原值，没有改变上一轮排除规则。",
           "原论文GF四事件数19、37、21、44，共121；原表非空GF日期编号数也为121。上一轮保守排除上述3个Staff描述编号后留下118个主/探索候选。数量吻合不能证明逐条身份和作者清洗规则一致，因此不宣布完成作者数据复现。", "",
           "## 保存预测的阶段诊断", "",
           "仅提取上一轮15条最早记录的输出，没有重新拟合或调用LLM。Logit来自上一轮含改道探索记录的按日期留出模型。", "",
           "| 方法 | 符合最初出口1 | 符合后来出口2 | 预测其他出口 |", "|---|---:|---:|---:|"]
    for r in dsummary:
        lines.append(f"| {r['method']} | {r['matches_initial']}/15 | {r['matches_later']}/15 | {r['other_exit']}/15 |")
    lines += ["", "这些输出只读取最早属性。“符合后来出口2”是描述性核查，不能改名为动态改道预测准确率。因为全部序列都选自已改道者，永远预测改到出口2就会取得15/15，完全不需要现场信息。",
              "同样，103条单条记录不能自动构造为收到同样提醒后仍不改道的负样本，它们没有匹配的第二次观测，也可能只是没有记录完整过程。", "",
              "## 对研究主线的决定", "",
              "当前停止把这15条当作“是否改道、何时改道”的正式训练/评价集，保留为描述和待核对序列。继续冻结原PPO扩大训练和48个最终疏散场景。",
              "现有记录能继续支持初次/单次选择预测。下一项可实施的有限比较是：在训练日期内给Qwen真人示例或进行概率校准，再和用相同训练日期的Logit比较，以检验真人学习机会能否改善其预测。该比较须另行冻结，属于探索，不重复包装成本批日期的最终确认结果。",
              "进入动态引导研究需要前后决策时点、提醒内容/方向及发生时间，以及同一决策窗口内改道者和保持者的可比观测。已生成15行复核表，未知时间和视频编号保持空白。", "",
              "## 输出与来源", "",
              "transition_pairs.json保存原属性及前后配对；transition_review.csv提供待核实字段；staff_role_audit.csv保存身份依据；cached_stage_diagnostic.csv保存逐人阶段诊断。",
              "原附件和上一轮清洗、冻结计划、全部221条回答及报告均进行前后哈希核对，没有改写前一轮结果。",
              "[Tong等，Safety Science 182 (2025), 106708](https://doi.org/10.1016/j.ssci.2024.106708)，方法3.3及表3解释最终Guide定义和样本筛选边界。用户附件Excel和PPT为本轮逐行核查来源。", ""]
    report=ROOT/"reports/Exit_choice改道序列与工作人员编码核查_2026-10-12.md"
    report.write_text("\n".join(lines),encoding="utf-8")
    after=[{**r,"sha256_after":sha(r["path"])} for r in before]
    for r in after:
        r["unchanged"] = r["sha256"]==r["sha256_after"]
    assert all(r["unchanged"] for r in after)
    save("preservation_after.json",after)
    save("completion_status.json",{"completed_at_utc":datetime.now(timezone.utc).isoformat(),
         "state":"completed_source_transition_audit","prior_research_files_unchanged":len(after),
         "report":str(report),"report_sha256":sha(report), **summary})
    print(json.dumps({"report":str(report),"protected_files_unchanged":len(after),**summary},ensure_ascii=False))


if __name__=="__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()

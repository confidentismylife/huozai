"""Create a concise, source-grounded five-minute research briefing."""
from pathlib import Path
import hashlib
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf as fitz

from build_report_docx import configure_styles, set_run_font, add_page_number
from build_full_history_report import inline
from build_integrated_report_documents_20261008 import export_pdf

ROOT=Path(__file__).resolve().parents[1]
NAME='LLM_RL火灾疏散研究_5分钟简洁版_2026-10-08'
SOURCE=ROOT/'reports/LLM_RL火灾疏散研究全程整合报告_2026-10-08.md'
ASSETS=ROOT/'artifacts/five_minute_report_assets_20261008'

CONTENT='''# LLM＋RL火灾疏散研究：5分钟简洁版

**阶段汇报｜数据截至2026年10月8日｜可直接朗读，约5分钟**

## 1. 我们想解决什么问题？（约45秒）

我们研究的是：火灾中，能不能让大语言模型提出疏散建议，再由强化学习根据现场变化调整，从而减少烟气暴露、拥堵和人员失能。

最初还希望用LLM模拟人的出口选择。后来考虑到真实行为验证不足、逐人调用成本高，主要控制实验转向“LLM提出高层建议，PPO负责动态执行”。这里的PPO，可以理解为在仿真中学会调节出口指示与放行量的控制器。

现在最关键的问题是：**面对一条建议，系统能否判断此刻采纳它比继续执行PPO更好？**

## 2. 我们已经做了什么？（约60秒）

我们已经搭建了可运行的简化疏散平台，包含两个房间、一个路口、两条出口走廊，以及烟气增长、出口封闭、拥堵和不同的人群行为。控制器可以调整出口方向和分级放行容量。

研究逐步完成了PPO训练、多训练种子比较、行为与观测噪声变化实验、本地LLM接入，以及建议和实际动作的逐步审计。相应轮次还补齐了相同动作、信息和训练预算的对照，避免把条件不同的成绩直接相比。

近期控制诊断使用Qwen3 14B已有响应缓存。系统采用“现场观测→LLM候选→规则与信任门控→PPO执行”的流程，检查建议是否过期、出口是否封闭，并限制放行容量。PPO也遵守同样的约束，仍需要单独验证效果。

## 3. 实验得到了什么结果？（约90秒）

**第一，尚未证明LLM＋PPO稳定优于纯PPO。**第一轮公平对照中，LLM方案的失能率高出1.34个百分点，区间为[+0.67，+2.10]。第二轮加入护栏后，平均仍高出1.02个百分点，但区间为[-1.57，+4.55]，跨过零，因此不能判断存在稳定改善。这些结果只针对各轮具体配置，也不能推断所有LLM都无效。

**第二，门控需要同时识别有害建议和有益建议。**最近完整实验完成2,080个仿真回合、2,480次反事实终点模拟。即时模式下，24个再次覆盖机会中，有13个有害、5个有益、6个收益微小或近于中性。两种门控全部拒绝：拦住了13个有害建议，也错过了全部5个有益建议。因此，“没有接受坏建议”还不能说明系统会判断收益。

这些回合覆盖的是五个同一简化图上的参数情景及不同人口、行为和权重，不能当成大量独立建筑。不同轮次条件不同，也不能把成绩连成持续提升的曲线。

## 4. 为什么还没有稳定提升？（约45秒）

已经发现，出口方向、放行容量和采纳时机相互影响。看起来更安全的出口，如果吞吐量低，也可能让人停留更久、暴露更多；第一次有用的建议，再次执行时可能变差。

另外，现有观测和公开行动历史相同的两个病例，候选收益仍一正一负；背后的人群组成与行为模型不同。这个发现提示我们需要检验响应观测能否增加有效信息，但还不能证明增加观测一定能解决问题。

当前主要控制实验使用合成烟气和合成人群参数，尚未完成真实建筑、物理烟气耦合和连续行人轨迹验证，结论应限定在现有模拟器内。

## 5. 接下来具体做什么？（约60秒）

下一步先研究“观察到的人群响应，能否帮助预测当前建议的收益”。例如，观察两侧走廊进入人数、排队消退速度，以及改变指示后人流怎样变化；这些指标需要验证，不能直接当作真实服从率。

先用已保存轨迹整理严格早于决策的历史，排除未来信息泄漏，比较仅用当前观测、加入普通历史、加入响应特征三种方案；分别预测失能变化与暴露变化，再冻结设置，用新人口验证。评价要同时统计有害拦截、有益误拒，以及完整回合的失能、暴露、疏散和拥堵结果。

这一阶段目前只有特征与轻量预测器代码初稿，尚未测试、拟合或验证。暂缓扩大PPO训练，48个保留最终情景继续封存。

**目前已完成的是实验平台、公平对照和失败机制诊断；下一阶段要补上“何时值得采纳建议”的预测证据。稳定控制优势和现实有效性仍待验证。**

---

证据索引（不计入口头汇报）：研究目标与系统见完整版第1—3章；公平对照见第7—8章；近期控制、限制与后续设计见第36—39章。本文只压缩已有材料，没有新增实验。

[完整报告与逐轮证据](C:/Users/摩天轮大盗/Documents/ChatGPT/LLM探索/reports/LLM_RL火灾疏散研究全程整合报告_2026-10-08.md)
'''


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    preserved=[SOURCE,ROOT/'artifacts/LLM_RL火灾疏散研究全程整合报告_2026-10-08.docx',ROOT/'artifacts/LLM_RL火灾疏散研究全程整合报告_2026-10-08.pdf',ROOT/'outputs/local_llm_selection.json',ROOT/'outputs/third_round_development_status.json']
    before={str(p):sha(p) for p in preserved}
    md=ROOT/'reports'/(NAME+'.md');docx=ROOT/'artifacts'/(NAME+'.docx');pdf=docx.with_suffix('.pdf')
    ASSETS.mkdir(parents=True,exist_ok=True)
    if any(p.exists() for p in (md,docx,pdf)):
        assert '--rebuild' in sys.argv
        if pdf.exists():pdf.replace(ASSETS/'previous_layout.pdf')
    md.write_text(CONTENT,encoding='utf-8')
    doc=Document();configure_styles(doc)
    section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.63)
    section.left_margin=section.right_margin=Inches(.73)
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.2
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.first_line_indent=Pt(21)
    h=doc.styles['Heading 1'];h.font.size=Pt(12)
    h.paragraph_format.space_before=Pt(10);h.paragraph_format.space_after=Pt(5)
    expected=[]
    for line in CONTENT.splitlines():
        if not line.strip() or line=='---':continue
        if line.startswith('# '):
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_after=Pt(8)
            set_run_font(p.add_run(line[2:]),chinese='微软雅黑',size=18,bold=True)
            expected.append(line[2:]);continue
        if line.startswith('## '):
            p=doc.add_heading(line[3:],1)
            if line.startswith('## 4.'):p.paragraph_format.page_break_before=True
            expected.append(line[3:]);continue
        p=doc.add_paragraph()
        if line.startswith('**阶段'):
            p.paragraph_format.first_line_indent=Pt(0);inline(p,line,9)
        elif line.startswith(('证据索引','[完整报告')):
            p.paragraph_format.first_line_indent=Pt(0);inline(p,line,8)
        else:inline(p,line,10.5)
        expected.append(re.sub(r'\*\*','',re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',line)))
    header=section.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run('LLM＋RL火灾疏散研究｜5分钟简洁版｜2026-10-08'),size=8,color='666666')
    add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title='LLM＋RL火灾疏散研究：5分钟简洁版'
    doc.core_properties.subject='基于全程整合报告的约5分钟阶段汇报'
    doc.save(docx);export_pdf(docx)
    normalize=lambda s:re.sub(r'\s+','',s)
    with zipfile.ZipFile(docx) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        word=''.join(n.text or '' for n in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    d=fitz.open(pdf);text='\n'.join(p.get_text() for p in d)
    assert all(normalize(s) in normalize(word) and normalize(s) in normalize(text) for s in expected)
    assert len(d)<=3 and '\ufffd' not in text
    for i,page in enumerate(d):
        assert page.get_text().strip()
        for b in page.get_text('blocks'):
            assert b[0]>=-1 and b[1]>=-1 and b[2]<=page.rect.width+1 and b[3]<=page.rect.height+1
        page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(ASSETS/f'page_{i+1:02d}.png')
    assert all(sha(Path(p))==h for p,h in before.items())
    main_text=CONTENT.split('---')[0]
    count=len(re.findall(r'[\u4e00-\u9fff]',main_text))
    qa=dict(pages=len(d),main_chinese_characters=count,estimated_minutes=[round(count/320,1),round(count/260,1)],expected_text_units=len(expected),missing_text=0,preservation_checked=before,source=str(SOURCE),new_experiments=0,new_llm_requests=0,reserved_final_scenarios_read=0,artifacts=[dict(path=str(p),sha256=sha(p),bytes=p.stat().st_size) for p in (md,docx,pdf)])
    (ASSETS/'qa_manifest.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in qa.items() if k not in ('preservation_checked','artifacts')},ensure_ascii=False),flush=True)


if __name__=='__main__':main()

"""Build a Chinese retrospective report without modifying experiment outputs."""
from __future__ import annotations

import argparse
from pathlib import Path, PurePosixPath, PureWindowsPath
import re
import unicodedata

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_report_docx import configure_styles, set_run_font, set_cell_shading, set_cell_margins, add_page_number

ROOT = Path(__file__).resolve().parents[1]
NAME = "LLM_RL火灾疏散全流程实验研究总报告_2026-10-02"


def hyperlink_target(target):
    """Use file URIs for absolute local paths, while preserving web targets."""
    target = target.strip()
    if target.startswith("<") and target.endswith(">"):
        target = target[1:-1]
    if re.match(r"^[A-Za-z]:[\\/]", target):
        # Codex file links can include a line suffix; Word opens the file itself.
        return PureWindowsPath(re.sub(r":\d+$", "", target)).as_uri()
    if target.startswith("/"):
        return PurePosixPath(re.sub(r":\d+$", "", target)).as_uri()
    return target


def hyperlink(paragraph, label, target, size=10.5, bold=False, italic=False):
    node = OxmlElement("w:hyperlink")
    relation = paragraph.part.relate_to(hyperlink_target(target), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    node.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245B78")
    properties.append(color)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(round(size * 2)))
    properties.append(font_size)
    east_asian_size = OxmlElement("w:szCs")
    east_asian_size.set(qn("w:val"), str(round(size * 2)))
    properties.append(east_asian_size)
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    node.append(run)
    paragraph._p.append(node)


def markdown_link(text, start):
    """Read a link with balanced target parentheses, including local paths."""
    if text[start:start + 1] != "[":
        return None
    label_end = start + 1
    while label_end < len(text):
        if text[label_end] == "\\":
            label_end += 2
            continue
        if text[label_end] == "]":
            break
        label_end += 1
    if text[label_end:label_end + 2] != "](":
        return None
    target_start = label_end + 2
    cursor, depth, angled = target_start, 1, text[target_start:target_start + 1] == "<"
    while cursor < len(text):
        char = text[cursor]
        if char == "\\" and text[cursor + 1:cursor + 2] in ("(", ")"):
            cursor += 2
            continue
        if angled:
            if char == ">" and text[cursor + 1:cursor + 2] == ")":
                cursor += 1
                depth = 0
                break
        elif char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
            if depth == 0:
                break
        cursor += 1
    if depth:
        return None
    label = text[start + 1:label_end]
    target = text[target_start:cursor].replace(r"\(", "(").replace(r"\)", ")")
    return label, target, cursor + 1


def inline(paragraph, text, size=10.5, bold=False, italic=False):
    """Render links, code and emphasis without losing text inside markup."""
    position, plain = 0, []

    def flush():
        if plain:
            run = paragraph.add_run("".join(plain))
            set_run_font(run, size=size, bold=bold)
            run.italic = italic
            plain.clear()

    while position < len(text):
        if (text[position] == "\\" and position + 1 < len(text) and
                text[position + 1] in "\\`*[]|"):
            plain.append(text[position + 1])
            position += 2
            continue
        if text.startswith("**", position):
            end = text.find("**", position + 2)
            if end >= 0:
                flush()
                inline(paragraph, text[position + 2:end], size=size, bold=True, italic=italic)
                position = end + 2
                continue
        if text[position] == "*" and not text.startswith("**", position):
            end = text.find("*", position + 1)
            if end > position + 1:
                flush()
                inline(paragraph, text[position + 1:end], size=size, bold=bold, italic=True)
                position = end + 1
                continue
        if text[position] == "`":
            marker = re.match(r"`+", text[position:])[0]
            end = text.find(marker, position + len(marker))
            if end >= 0:
                flush()
                run = paragraph.add_run(text[position + len(marker):end])
                set_run_font(run, chinese="等线", latin="Consolas", size=max(8, size-1.5), bold=bold)
                run.italic = italic
                position = end + len(marker)
                continue
        link = markdown_link(text, position)
        if link:
            flush()
            hyperlink(paragraph, link[0], link[1], size=size, bold=bold, italic=italic)
            position = link[2]
            continue
        plain.append(text[position])
        position += 1
    flush()


def split_table_row(text):
    """Do not treat escaped or inline-code pipes as column boundaries."""
    text = text.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|") and not text.endswith(r"\|"):
        text = text[:-1]
    cells, current, code_marker = [], [], None
    position = 0
    while position < len(text):
        char = text[position]
        if char == "\\" and text[position + 1:position + 2] == "|":
            current.append(r"\|")
            position += 2
            continue
        if char == "`":
            marker = re.match(r"`+", text[position:])[0]
            if code_marker is None:
                code_marker = marker
            elif marker == code_marker:
                code_marker = None
            current.append(marker)
            position += len(marker)
            continue
        if char == "|" and code_marker is None:
            cells.append("".join(current).strip())
            current.clear()
        else:
            current.append(char)
        position += 1
    cells.append("".join(current).strip())
    return cells


def visible_length(text):
    """Approximate rendered column text, rather than counting hidden URLs."""
    position, visible = 0, []
    while position < len(text):
        link = markdown_link(text, position)
        if link:
            visible.append(link[0])
            position = link[2]
        else:
            visible.append(text[position])
            position += 1
    plain = "".join(visible).replace("*", "").replace("`", "")
    return sum(1 if unicodedata.east_asian_width(char) in ("F", "W") else .6 for char in plain)


def table(doc, rows, available_width):
    if not rows:
        raise ValueError("Markdown table has no content rows")
    cols = len(rows[0])
    if not all(len(row) == cols for row in rows):
        raise ValueError("ragged Markdown table")
    obj = doc.add_table(rows=0, cols=cols)
    obj.alignment = WD_TABLE_ALIGNMENT.CENTER
    obj.autofit = False
    obj.style = "Table Grid"
    lengths = [max(6, min(32, max(visible_length(row[c]) for row in rows))) for c in range(cols)]
    widths = [available_width * length / sum(lengths) for length in lengths]
    for col, width in zip(obj.columns, widths):
        col.width = Inches(width)
    for i, values in enumerate(rows):
        row = obj.add_row()
        props = row._tr.get_or_add_trPr()
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            props.append(repeat)
        cant_split = OxmlElement("w:cantSplit")
        props.append(cant_split)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            set_cell_margins(cell, top=70, bottom=70, start=80, end=80)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, "E7EFF4")
            elif i % 2 == 0:
                set_cell_shading(cell, "F6F8FA")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            inline(p, value, 9, bold=i == 0)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)


def ordered_numbering(doc, start):
    """Create a real Word list that restarts at its Markdown start number."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids, default=-1) + 1, max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for name, value in (("start", start), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")):
        element = OxmlElement(f"w:{name}")
        element.set(qn("w:val"), str(value))
        level.append(element)
    properties = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "360")
    indentation.set(qn("w:hanging"), "180")
    properties.append(indentation)
    level.append(properties)
    abstract.append(level)
    numbering.append(abstract)
    instance = OxmlElement("w:num")
    instance.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    instance.append(reference)
    numbering.append(instance)
    return num_id


def set_orientation(doc, landscape):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Inches(11.69), Inches(8.27)) if landscape else (Inches(8.27), Inches(11.69))
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.75)
    return float((section.page_width - section.left_margin - section.right_margin) / Inches(1))


def build(source, output):
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.75)
    available = 6.77
    configure_styles(doc)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Normal"].paragraph_format.widow_control = True
    for text, size in (("LLM与强化学习火灾疏散研究", 22), ("全流程实验与方法复盘报告", 19)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(40 if size == 22 else 14)
        p.paragraph_format.first_line_indent = Pt(0)
        set_run_font(p.add_run(text), chinese="微软雅黑", size=size, bold=True)
    for text in ("从研究设想、代码原型与真实数据验证", "到动态门控、公平对照、护栏审计与计划蒸馏", "", "截至2026年10月2日的阶段性证据", "用途：导师讨论、实验复盘、后续研究设计"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        set_run_font(p.add_run(text), chinese="微软雅黑", size=12, color="444444")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(35)
    inline(p, "本报告完整保留正面与负面实验。仿真链路已实现，当前主方法尚未建立稳定LLM增益或通过硬护栏资格。控制实验仍以合成烟气和人员模型为主。")
    doc.add_page_break()
    doc.add_heading("报告导航", 1)
    for line in lines:
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            inline(p, line[3:], 11)
    doc.add_page_break()
    i = 0
    figures = []
    ordered_num_id, previous_number, wide_tables = None, None, 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("# "):
            i += 1
            continue
        ordered = re.match(r"^(\d+)[.)]\s+(.+)$", line)
        if not ordered:
            ordered_num_id, previous_number = None, None
        heading = re.match(r"^(#{2,6}) (.+)$", line)
        if heading:
            doc.add_heading(heading[2], min(3, len(heading[1])-1))
            i += 1
            continue
        figure = re.fullmatch(r"!\[([^\]]*)\]\((.+)\)", line)
        if figure:
            path = Path(figure[2].strip("<>"))
            if not path.is_file():
                raise FileNotFoundError(path)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.add_run().add_picture(str(path), width=Inches(available))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            inline(p, figure[1], 9)
            figures.append(str(path))
            i += 1
            continue
        fence = re.match(r"^(`{3,}|~{3,})", line)
        if fence:
            marker = fence[1]
            closing = re.compile(re.escape(marker[0]) + "{" + str(len(marker)) + r",}\s*$")
            i += 1
            while i < len(lines) and not closing.fullmatch(lines[i].strip()):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                set_run_font(p.add_run(lines[i]), chinese="等线", latin="Consolas", size=8)
                i += 1
            if i >= len(lines):
                raise ValueError("unclosed Markdown code fence")
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = split_table_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in row):
                    rows.append(row)
                i += 1
            if rows and len(rows[0]) >= 8:
                wide_width = set_orientation(doc, landscape=True)
                table(doc, rows, wide_width)
                set_orientation(doc, landscape=False)
                wide_tables += 1
            else:
                table(doc, rows, available)
            continue
        bullet = re.match(r"^[-+*]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Pt(0)
            inline(p, bullet[1])
            i += 1
            continue
        if ordered:
            number = int(ordered[1])
            if ordered_num_id is None or number != previous_number + 1:
                ordered_num_id = ordered_numbering(doc, number)
            p = doc.add_paragraph(style="List Number")
            properties = p._p.get_or_add_pPr().get_or_add_numPr()
            properties.get_or_add_ilvl().val = 0
            properties.get_or_add_numId().val = ordered_num_id
            inline(p, ordered[2])
            previous_number = number
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph(style="Quote")
            inline(p, line[2:])
            i += 1
            continue
        if re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", line):
            i += 1
            continue
        p = doc.add_paragraph()
        inline(p, line)
        i += 1
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("LLM＋RL火灾疏散｜全流程实验复盘｜2026-10-02"), chinese="微软雅黑", size=8, color="777777")
    add_page_number(section.footer.paragraphs[0])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "LLM与强化学习火灾疏散全流程实验研究总报告"
    doc.core_properties.subject = "既有实验历史、方法复盘与证据边界"
    doc.core_properties.author = "项目研究记录"
    doc.save(output)
    return {"output": str(output), "paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
            "wide_tables": wide_tables, "figures": len(figures), "sections": len(doc.sections)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=ROOT / "reports" / f"{NAME}.md")
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts" / f"{NAME}.docx")
    args = parser.parse_args()
    print(build(args.source, args.output))

"""Auditable, bounded offline human-choice benchmark; no evacuation/PPO files read.

Commands: prepare, baselines, llm, report. Sources are read-only. Each inference
request and response is saved separately; resuming never repeats a saved request.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import time
import urllib.request

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/exit_choice_initial_benchmark_20261012"
SOURCES = [Path("C:/Users/摩天轮大盗/Downloads/Exit choice.xlsx"),
           Path("C:/Users/摩天轮大盗/Downloads/Exit choice ppt.pptx")]
PROTECTED = [ROOT / "outputs/local_llm_selection.json",
             ROOT / "outputs/third_round_development_status.json"]
BASE = "http://127.0.0.1:11434"
DATES = ["2019-03-02", "2019-03-06", "2019-09-11", "2019-10-31"]
# Verified against raw header cells, exit names, and PPT slides 3/5/10/15.
# The former auto-extractor missed headings without the literal word 'Level'.
BLOCKS = {
    "02-03-19 EC": {2: "general", 8: "general", 16: "cafe_corridor"},
    "06-03-19 EC": {74: "cafe_corridor", 97: "computer_lab_east",
                      109: "computer_lab_west", 128: "lending_desk_1",
                      133: "lending_desk_2"},
    "11-09-19 EC": {3: "computer_lab_east", 11: "computer_lab_west",
                      218: "lending_desk_1"},
    "31-10-19 EC": {84: "cafe_corridor", 114: "computer_lab_west",
                      129: "computer_lab_east"},
}
SLIDES = dict(zip(BLOCKS, [3, 5, 10, 15]))
METHODS = ["uniform", "training_frequency", "nearest_exit", "logit"]


def now():
    return datetime.now(timezone.utc).isoformat()


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(name, value):
    p = OUT / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False,
                            default=str), encoding="utf-8")


def read(name):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def write_csv(name, rows):
    rows = list(rows)
    with (OUT / name).open("w", encoding="utf-8-sig", newline="") as f:
        if rows:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]))
            writer.writeheader()
            for row in rows:
                writer.writerow({k: json.dumps(v, ensure_ascii=False) if
                                 isinstance(v, (dict, list)) else v
                                 for k, v in row.items()})


def api(path, body=None, timeout=10):
    req = urllib.request.Request(BASE + path,
                                 data=None if body is None else json.dumps(body).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as f:
        return json.load(f)


def check_preservation():
    p = read("preservation_before.json")
    results = [{**r, "sha256_after": sha(r["path"])} for r in p]
    for r in results:
        r["unchanged"] = r["sha256"] == r["sha256_after"]
    save("preservation_after.json", results)
    assert all(r["unchanged"] for r in results), "Protected source changed"


def canonical_id(v):
    if isinstance(v, (int, float)) and int(v) == v:
        return str(int(v))
    return str(v).strip()


def choice(v):
    if isinstance(v, (int, float)) and v in (1, 2, 3):
        return int(v), int(v), "single_numeric"
    if isinstance(v, str):
        m = re.fullmatch(r"\s*([123])\s+then\s+([123])(?:\s+due\s+to\s+staff)?\s*", v, re.I)
        if m:
            return int(m[1]), int(m[2]), "sequence_annotation"
    return None, None, "missing_or_unparsed"


def prepare():
    if (OUT / "frozen_plan.json").exists():
        check_preservation()
        print("Existing preparation preserved; no plan overwritten", flush=True)
        return
    OUT.mkdir(parents=True, exist_ok=True)
    save("preservation_before.json", [{"path": str(p), "sha256": sha(p)}
                                      for p in SOURCES + PROTECTED if p.exists()])
    wb = openpyxl.load_workbook(SOURCES[0], read_only=True, data_only=False)
    raw, block_manifest = [], []
    for sheet, blocks in BLOCKS.items():
        ws = wb[sheet]
        all_rows = list(ws.iter_rows(values_only=True))
        active = None
        for rownum, vals in enumerate(all_rows, 1):
            starts = [i for i, v in enumerate(vals) if isinstance(v, str) and v.strip().lower() == "dist"]
            if starts:
                active = None
                if rownum in blocks:
                    assert len(starts) == 3
                    ch = [i for i, v in enumerate(vals) if isinstance(v, str) and v.strip().lower() == "choice"]
                    assert len(ch) == 1
                    # Header names provide independent physical-exit correspondence.
                    exit_labels = [all_rows[rownum-2][i] for i in starts]
                    assert "main entrance" in str(exit_labels[0]).lower()
                    active = {"starts": starts, "choice": ch[0], "header_row": rownum,
                              "region": blocks[rownum], "exit_labels": exit_labels}
                    block_manifest.append({"sheet": sheet, **active,
                                           "ppt_slide": SLIDES[sheet], "floor": "Ground Floor / Level 2"})
                continue
            if not active or vals[0] is None or not re.fullmatch(r"\d+[a-zA-Z]?", canonical_id(vals[0])):
                continue
            ident = canonical_id(vals[0])
            date = datetime.strptime(sheet[:8], "%d-%m-%y").date().isoformat()
            initial, final, kind = choice(vals[active["choice"]])
            desc = " ".join(str(v) for v in vals[1:active["starts"][0]] if v is not None)
            attrs = [{"exit": j+1, **dict(zip(["dist", "people", "staff", "fam"],
                                              vals[start:start+4]))}
                     for j, start in enumerate(active["starts"])]
            raw.append({"record_id": f"{date}:r{rownum}", "date": date, "sheet": sheet,
                        "row": rownum, "person_id": ident, "person_key": f"{date}:{ident}",
                        "region": active["region"], "header_row": active["header_row"],
                        "choice_cell": f"{openpyxl.utils.get_column_letter(active['choice']+1)}{rownum}",
                        "raw_choice": vals[active["choice"]], "initial_label": initial,
                        "final_label": final, "label_kind": kind, "exits": attrs,
                        "description_for_matching_only": desc,
                        "staff_role": bool(re.search(r"\bstaff\b", desc, re.I)),
                        "description_is_model_input": False,
                        "ppt_slide": SLIDES[sheet], "input_timing": "raw_record_not_video_verified"})
    wb.close()
    grouped = defaultdict(list)
    for r in raw:
        grouped[r["person_key"]].append(r)
    strict, sequences, row_audit, exclusions = [], [], [], []
    for person, records in grouped.items():
        records.sort(key=lambda r: r["row"])
        first = records[0]
        reasons = []
        if any(r["staff_role"] for r in records):
            reasons.append("staff_role")
        if first["initial_label"] is None:
            reasons.append("missing_or_unparsed_choice")
        for field in ["dist", "people", "staff", "fam"]:
            values = [e[field] for e in first["exits"]]
            if any(not isinstance(v, (int, float)) or not math.isfinite(v) or v < 0 for v in values):
                reasons.append(f"missing_nonfinite_or_censored_{field}")
            elif field in ("staff", "fam") and any(v not in (0, 1) for v in values):
                reasons.append(f"nonbinary_{field}_unresolved")
        changed = len(records) > 1 or first["label_kind"] == "sequence_annotation"
        if changed:
            reasons.append("sequence_input_timing_unverified")
        disposition = "exclude"
        if not reasons:
            disposition = "strict_single_record"
            strict.append(first)
        elif reasons == ["sequence_input_timing_unverified"]:
            disposition = "exploratory_first_sequence_record"
            first = {**first, "input_timing": "provisional_earliest_row_before_change_not_verified"}
            sequences.append(first)
        if reasons:
            exclusions.append({"person_key": person, "date": first["date"],
                               "source_rows": [r["row"] for r in records],
                               "disposition": disposition, "reasons": reasons})
        for j, r in enumerate(records):
            row_audit.append({"record_id": r["record_id"], "date": r["date"],
                              "sheet": r["sheet"], "row": r["row"], "person_id": r["person_id"],
                              "choice_cell": r["choice_cell"], "raw_choice": r["raw_choice"],
                              "region": r["region"], "person_record_count": len(records),
                              "disposition": disposition if j == 0 else "later_record_not_independent_sample",
                              "reasons": reasons})
    strict.sort(key=lambda r: r["record_id"])
    sequences.sort(key=lambda r: r["record_id"])
    # No claim that a suffix identifies the same person: 17b is explicitly a friend.
    # PPT IDs are used only for matching; arrows/colours are never model inputs.
    source_summary = read_external_inspection()
    ppt_audit = []
    for sheet, slide in SLIDES.items():
        ppt_ids = {s.strip() for s in source_summary[slide-1]["texts"] if re.fullmatch(r"\d+[a-zA-Z]?", s.strip())}
        raw_ids = {r["person_id"] for r in raw if r["sheet"] == sheet}
        ppt_audit.append({"sheet": sheet, "slide": slide, "raw_ids": len(raw_ids),
                          "ppt_ids": len(ppt_ids), "in_excel_not_ppt": sorted(raw_ids-ppt_ids),
                          "in_ppt_not_excel": sorted(ppt_ids-raw_ids)})
    save("block_mapping.json", block_manifest)
    save("raw_level2_records.json", raw)
    save("strict_initial_cases.json", strict)
    save("exploratory_sequence_cases.json", sequences)
    save("ppt_matching_audit.json", ppt_audit)
    write_csv("row_audit.csv", row_audit)
    write_csv("exclusion_reasons.csv", exclusions)
    case_rows = []
    for r in strict + sequences:
        for e in r["exits"]:
            case_rows.append({"case_id": r["record_id"], "date": r["date"], "person_key": r["person_key"],
                              "sheet": r["sheet"], "source_row": r["row"], "region": r["region"],
                              "cohort": "strict" if r in strict else "sequence_exploratory",
                              "initial_choice_label": r["initial_label"], **e,
                              "input_timing": r["input_timing"]})
    write_csv("clean_choice_long.csv", case_rows)
    summary = {
        "raw_rows": len(raw), "person_keys": len(grouped), "strict_cases": len(strict),
        "sequence_exploratory_cases": len(sequences),
        "later_rows": sum(len(v)-1 for v in grouped.values()),
        "strict_per_date": dict(Counter(r["date"] for r in strict)),
        "sequence_per_date": dict(Counter(r["date"] for r in sequences)),
        "exclusion_reasons_person_level": dict(Counter(x for r in exclusions for x in r["reasons"])),
        "staff_values": dict(Counter(str(e["staff"]) for r in raw for e in r["exits"])),
        "fam_patterns_strict": dict(Counter(str([e["fam"] for e in r["exits"]]) for r in strict)),
        "ppt_matching": ppt_audit,
        "not_a_replication_of_paper_final_497": True,
    }
    save("data_summary.json", summary)
    # Frozen before fitting models or obtaining any real-case LLM answers.
    plan = {
        "frozen_at_utc": now(), "task": "offline recorded initial human exit choice",
        "primary_cohort": "strict_single_record_only",
        "exploratory_cohort": "strict_plus_provisional_first_sequence_row",
        "primary_channel": ["dist", "people", "fam"],
        "conditional_guide_channel": ["dist", "people", "fam", "staff"],
        "guide_channel_limit": "Raw Guide timing not individually video verified; no causal interpretation",
        "llm_model": "qwen3:8b", "llm_model_selection_source": str(PROTECTED[0]),
        "llm_system_prompt": SYSTEM, "llm_schema": SCHEMA,
        "source_script_sha256": sha(__file__),
        "llm_zero_shot": True, "llm_examples": 0,
        "llm_options": {"temperature": 0, "seed": 2026101201, "num_ctx": 2048,
                        "num_predict": 128, "num_batch": 32, "num_thread": 8, "num_gpu": 0},
        "llm_timeout_seconds": 180, "llm_retries": 0, "stop_after_consecutive_failures": 3,
        "llm_probability_sum_tolerance": 0.01,
        "outer_validation": "leave_one_date_out", "dates": DATES,
        "inner_validation": "leave_one_training_date_out",
        "logit_ridge_sum_nll_grid": [0.01, 0.1, 1.0, 10.0],
        "nearest_temperature_grid_raw_dist_units": [1.0, 2.0, 5.0, 10.0, 20.0, 50.0],
        "frequency_laplace_alpha": 1.0,
        "logit_features": "ASC_E1, ASC_E2, dist/10, people/10; optional binary Guide",
        "fam_handling": "Fam=[1,0,0] is collinear with ASC_E1; no independent Fam coefficient",
        "forbidden_llm_inputs": ["actual_choice", "final_choice", "change_notes", "date", "event_type",
                                 "person_id", "appearance", "PPT_arrows", "PPT_colours"],
        "metrics": ["log_loss", "multiclass_brier_sum", "accuracy", "exit_share_total_variation"],
        "log_loss_probability_floor": 1e-6,
        "invalid_llm_output_fallback": "outer_training_frequency; reported separately",
        "success_gates": "No automatic PPO entry from this exploratory benchmark; examine per-date errors and effect size",
        "frozen_test_scenarios_accessed": False,
        "prior_labels_inspected": True,
        "case_sha256": {name: sha(OUT/name) for name in
                        ["strict_initial_cases.json", "exploratory_sequence_cases.json"]},
    }
    save("frozen_plan.json", plan)
    (OUT/"source_snapshot").mkdir(exist_ok=True)
    (OUT/"source_snapshot/exit_choice_initial_benchmark.py").write_bytes(Path(__file__).read_bytes())
    check_preservation()
    print(json.dumps(summary, ensure_ascii=False), flush=True)


def read_external_inspection():
    p = ROOT / "outputs/exit_choice_dataset_inspection_20261011/source_inspection.json"
    return json.loads(p.read_text(encoding="utf-8"))["slides"]


def softmax(x):
    z = x - x.max(axis=-1, keepdims=True)
    e = np.exp(z)
    return e/e.sum(axis=-1, keepdims=True)


def design(cases, guide):
    a = np.zeros((len(cases), 3, 5 if guide else 4))
    a[:, 0, 0] = 1
    a[:, 1, 1] = 1
    for i, r in enumerate(cases):
        for j, e in enumerate(r["exits"]):
            a[i, j, 2:4] = [e["dist"]/10, e["people"]/10]
            if guide:
                a[i, j, 4] = e["staff"]
    return a


def targets(cases):
    return np.array([r["initial_label"]-1 for r in cases], dtype=int)


def fit_logit(cases, guide, ridge):
    x, y = design(cases, guide), targets(cases)
    theta = np.zeros(x.shape[-1])
    def terms(t, need_hessian=True):
        z = np.einsum("njd,d->nj", x, t)
        p = softmax(z)
        zm = z.max(axis=1)
        loss = np.sum(zm + np.log(np.exp(z-zm[:, None]).sum(axis=1))-z[np.arange(len(y)), y]) + ridge*np.dot(t,t)/2
        ex = np.einsum("nj,njd->nd", p, x)
        grad = (ex-x[np.arange(len(y)), y]).sum(axis=0)+ridge*t
        if not need_hessian:
            return loss
        hess = np.einsum("nj,njd,nje->de", p, x, x)-np.einsum("nd,ne->de", ex, ex)+ridge*np.eye(len(t))
        return loss, grad, hess
    converged = False
    for it in range(150):
        loss, grad, hess = terms(theta)
        if np.linalg.norm(grad, ord=np.inf) < 1e-7:
            converged = True
            break
        step = np.linalg.solve(hess, grad)
        scale = 1.
        while scale > 1e-10 and terms(theta-scale*step, False) > loss-1e-4*scale*np.dot(grad, step):
            scale /= 2
        theta -= scale*step
    loss, grad, _ = terms(theta)
    if np.linalg.norm(grad, ord=np.inf) < 1e-6:
        converged = True
    assert converged, f"Logit optimizer failed, gradient={grad}"
    return theta, {"ridge": ridge, "iterations": it+1, "converged": converged,
                   "max_abs_gradient": float(abs(grad).max()), "penalized_sum_nll": float(loss),
                   "coefficients": theta.tolist(),
                   "coefficient_names": ["ASC_E1_includes_Fam_proxy", "ASC_E2", "dist_per10", "people_per10"] + (["Guide"] if guide else [])}


def metrics(cases, p):
    y = targets(cases)
    assert p.shape == (len(cases), 3)
    assert np.isfinite(p).all() and (p >= 0).all() and np.allclose(p.sum(axis=1), 1)
    truth = np.eye(3)[y]
    shares = truth.mean(axis=0)
    predicted = p.mean(axis=0)
    return {"n": len(y), "accuracy": float(np.mean(p.argmax(axis=1) == y)),
            "log_loss": float(-np.log(np.maximum(p[np.arange(len(y)), y], 1e-6)).mean()),
            "brier": float(((p-truth)**2).sum(axis=1).mean()),
            "exit_share_tv": float(abs(predicted-shares).sum()/2),
            "true_shares": shares.tolist(), "predicted_shares": predicted.tolist()}


def frequency(cases):
    return (np.bincount(targets(cases), minlength=3)+1)/(len(cases)+3)


def choose_parameter(train, grid, predictor):
    candidates = []
    for value in grid:
        folds = []
        for d in sorted({r["date"] for r in train}):
            tr = [r for r in train if r["date"] != d]
            va = [r for r in train if r["date"] == d]
            p = predictor(tr, va, value)
            folds.append({"date": d, "log_loss": metrics(va, p)["log_loss"]})
        candidates.append({"value": value, "mean_date_log_loss": float(np.mean([f["log_loss"] for f in folds])), "folds": folds})
    best = min(candidates, key=lambda c: (c["mean_date_log_loss"], c["value"]))
    return best["value"], candidates


def baselines():
    plan = read("frozen_plan.json")
    for name, digest in plan["case_sha256"].items():
        assert sha(OUT/name) == digest
    strict, seq = read("strict_initial_cases.json"), read("exploratory_sequence_cases.json")
    predictions, fit_details = [], []
    for cohort, cases in [("strict", strict), ("sequence_exploratory", strict+seq)]:
        for channel, guide in [("primary", False), ("guide_conditional", True)]:
            if cohort == "sequence_exploratory" and guide:
                continue
            for date in DATES:
                train = [r for r in cases if r["date"] != date]
                test = [r for r in cases if r["date"] == date]
                assert train and test
                assert not ({r["person_key"] for r in train} & {r["person_key"] for r in test})
                tau, tau_cv = choose_parameter(train, plan["nearest_temperature_grid_raw_dist_units"],
                    lambda tr, va, t: softmax(-np.array([[e["dist"] for e in r["exits"]] for r in va])/t))
                ridge, ridge_cv = choose_parameter(train, plan["logit_ridge_sum_nll_grid"],
                    lambda tr, va, v: softmax(np.einsum("njd,d->nj", design(va, guide), fit_logit(tr, guide, v)[0])))
                theta, fit = fit_logit(train, guide, ridge)
                fit_details.append({"cohort": cohort, "channel": channel, "test_date": date,
                                    "train_n": len(train), "test_n": len(test), "train_dates": sorted({r["date"] for r in train}),
                                    "nearest_tau": tau, "nearest_inner_cv": tau_cv,
                                    "logit_inner_cv": ridge_cv, **fit})
                probs = {
                    "uniform": np.full((len(test),3), 1/3),
                    "training_frequency": np.tile(frequency(train), (len(test),1)),
                    "nearest_exit": softmax(-np.array([[e["dist"] for e in r["exits"]] for r in test])/tau),
                    "logit": softmax(np.einsum("njd,d->nj", design(test, guide), theta)),
                }
                for method, pp in probs.items():
                    for r, p in zip(test, pp):
                        predictions.append({"cohort": cohort, "channel": channel, "test_date": date,
                                            "case_id": r["record_id"], "person_key": r["person_key"],
                                            "method": method, "true_exit": r["initial_label"],
                                            "probabilities": p.tolist(), "fallback": False})
    save("baseline_predictions.json", predictions)
    write_csv("baseline_predictions.csv", predictions)
    save("baseline_fit_details.json", fit_details)
    save("baseline_metrics.json", aggregate(predictions, strict, seq))
    check_preservation()
    print(json.dumps([r for r in read("baseline_metrics.json") if r["test_date"]=="pooled"], ensure_ascii=False), flush=True)


def aggregate(predictions, strict, seq):
    lookup = {r["record_id"]: r for r in strict+seq}
    groups = defaultdict(list)
    for r in predictions:
        for date in [r["test_date"], "pooled"]:
            groups[(r["cohort"], r["channel"], r["method"], date)].append(r)
    result = []
    for (cohort, channel, method, date), records in sorted(groups.items()):
        pp = np.array([r["probabilities"] for r in records])
        result.append({"cohort": cohort, "channel": channel, "method": method, "test_date": date,
                       **metrics([lookup[r["case_id"]] for r in records], pp),
                       "fallbacks": sum(r.get("fallback", False) for r in records)})
    return result


SCHEMA = {"type": "object", "properties": {f"p{i}": {"type": "number", "minimum": 0, "maximum": 1} for i in (1,2,3)},
          "required": ["p1", "p2", "p3"], "additionalProperties": False}
SYSTEM = (
    "预测图书馆疏散者的实际初次出口选择概率。现场报警要求立即从最近出口离开。"
    "输入是记录中的三个备选出口属性：distance是近似欧氏距离（不是完整走路距离）；"
    "previous_choosers是在此人开始朝出口移动时已选择该出口的其他人数（不是排队人数）；"
    "commonly_used是日常使用出口的熟悉程度代理（不是逐人问卷）。"
    "如有staff_guidance，它表示记录中工作人员是否向该出口引导，不能假设所有人服从。"
    "预测人的行为，包括可能的熟悉偏好和跟随人群。输出JSON对象p1、p2、p3，分别对应E1、E2、E3，"
    "各概率0到1且总和等于1。只输出JSON。"
)


def llm():
    plan = read("frozen_plan.json")
    assert SYSTEM == plan["llm_system_prompt"] and SCHEMA == plan["llm_schema"]
    strict, seq = read("strict_initial_cases.json"), read("exploratory_sequence_cases.json")
    tags = api("/api/tags")
    tag = next(m for m in tags["models"] if m["name"] == plan["llm_model"])
    selected = json.loads(PROTECTED[0].read_text(encoding="utf-8"))
    assert tag["digest"] == selected["digest"] and selected["model"] == plan["llm_model"]
    runtime = {"at_utc": now(), "tag": tag, "version": api("/api/version"),
               "system_prompt": SYSTEM, "schema": SCHEMA, "plan_sha256": sha(OUT/"frozen_plan.json")}
    if (OUT/"llm_runtime_frozen.json").exists():
        old = read("llm_runtime_frozen.json")
        assert old["tag"]["digest"] == tag["digest"] and old["system_prompt"] == SYSTEM
        assert old["schema"] == SCHEMA and old["plan_sha256"] == runtime["plan_sha256"]
    else:
        save("llm_runtime_frozen.json", runtime)
    tasks = [(r, "primary") for r in strict+seq] + [(r, "guide_conditional") for r in strict]
    # Deterministic hash order avoids running all records from one date first.
    tasks.sort(key=lambda t: hashlib.sha256((t[0]["record_id"]+t[1]).encode()).hexdigest())
    failure_streak = 0
    for index, (case, channel) in enumerate(tasks, 1):
        key = case["record_id"].replace(":", "_")+"_"+channel
        name = f"responses/{key}.json"
        if (OUT/name).exists():
            continue
        attrs = []
        for e in case["exits"]:
            value = {"exit": f"E{e['exit']}", "distance": e["dist"],
                     "previous_choosers": e["people"], "commonly_used": e["fam"]}
            if channel == "guide_conditional":
                value["staff_guidance"] = e["staff"]
            attrs.append(value)
        payload = {"model": plan["llm_model"], "messages": [
                   {"role": "system", "content": SYSTEM},
                   {"role": "user", "content": json.dumps({"exits": attrs}, ensure_ascii=False)}],
                   "format": SCHEMA, "think": False, "stream": False,
                   "keep_alive": "2m", "options": plan["llm_options"]}
        record = {"case_id": case["record_id"], "channel": channel, "started_at_utc": now(), "payload": payload}
        start = time.monotonic()
        try:
            response = api("/api/chat", payload, plan["llm_timeout_seconds"])
            record["response"] = response
            parsed = json.loads(response["message"]["content"])
            assert set(parsed) == {"p1", "p2", "p3"}, "Unexpected probability keys"
            assert all(type(parsed[k]) in (int, float) for k in parsed), "Invalid probability type"
            probs = np.array([parsed[f"p{i}"] for i in (1,2,3)], dtype=float)
            assert np.isfinite(probs).all() and ((0 <= probs)&(probs <= 1)).all(), "Probability bounds"
            assert abs(probs.sum()-1) <= plan["llm_probability_sum_tolerance"], "Probability sum"
            record.update({"valid": True, "raw_probability_sum": float(probs.sum()),
                           "probabilities": (probs/probs.sum()).tolist()})
            failure_streak = 0
        except Exception as exc:
            record.update({"valid": False, "error": f"{type(exc).__name__}: {exc}"})
            failure_streak += 1
        record["wall_seconds"] = time.monotonic()-start
        save(name, record)
        save("llm_progress.json", {"at_utc": now(), "tasks": len(tasks),
                                    "saved": len(list((OUT/"responses").glob("*.json"))),
                                    "last_case": case["record_id"], "last_channel": channel,
                                    "last_valid": record["valid"], "failure_streak": failure_streak})
        print(json.dumps({"index": index, "total": len(tasks), "case": case["record_id"],
                          "channel": channel, "valid": record["valid"],
                          "seconds": round(record["wall_seconds"],2)}, ensure_ascii=False), flush=True)
        if failure_streak >= plan["stop_after_consecutive_failures"]:
            print("Stopping after three consecutive failures; no silent model change", flush=True)
            break
    check_preservation()


def report():
    strict, seq = read("strict_initial_cases.json"), read("exploratory_sequence_cases.json")
    pred = read("baseline_predictions.json")
    lookup = {(r["case_id"],r["channel"]): r for r in
              [json.loads(p.read_text(encoding="utf-8")) for p in (OUT/"responses").glob("*.json")]} if (OUT/"responses").exists() else {}
    for cohort, cases, channel in [("strict", strict, "primary"), ("strict", strict, "guide_conditional"),
                                   ("sequence_exploratory", strict+seq, "primary")]:
        for case in cases:
            response = lookup.get((case["record_id"],channel))
            if response is None:
                continue
            train = [r for r in cases if r["date"] != case["date"]]
            pred.append({"cohort": cohort, "channel": channel, "test_date": case["date"],
                         "case_id": case["record_id"], "person_key": case["person_key"],
                         "method": "qwen3_8b", "true_exit": case["initial_label"],
                         "probabilities": response["probabilities"] if response["valid"] else frequency(train).tolist(),
                         "fallback": not response["valid"]})
    aggregate_metrics = aggregate(pred, strict, seq)
    save("all_predictions.json", pred)
    write_csv("all_predictions.csv", pred)
    save("all_metrics.json", aggregate_metrics)
    write_csv("all_metrics.csv", aggregate_metrics)
    summary = read("data_summary.json")
    fit = read("baseline_fit_details.json")
    save("report_implementation_revision.json", {
        "source_script_sha256": sha(__file__),
        "frozen_source_script_sha256": read("frozen_plan.json")["source_script_sha256"],
        "revision_scope": "Report adds data limitations, omits misleading uniform top1, and displays date-macro metrics. Predictions and frozen inference/statistical settings unchanged.",
    })
    llm_records = list(lookup.values())
    expected_requests = len(strict)*2+len(seq)
    runtime_summary = {"expected_requests": expected_requests, "saved_requests": len(llm_records),
                       "valid_requests": sum(r["valid"] for r in llm_records),
                       "invalid_requests": sum(not r["valid"] for r in llm_records),
                       "complete": len(llm_records) == expected_requests,
                       "median_wall_seconds": float(np.median([r["wall_seconds"] for r in llm_records])) if llm_records else None,
                       "total_wall_seconds_sum": sum(r["wall_seconds"] for r in llm_records)}
    save("runtime_summary.json", runtime_summary)
    # Paired raw per-person differences; descriptive only, four event clusters.
    paired = []
    for channel in ("primary", "guide_conditional"):
        model_records = {r["case_id"]: r for r in pred if r["cohort"]=="strict" and r["channel"]==channel and r["method"]=="qwen3_8b"}
        baseline_records = {r["case_id"]: r for r in pred if r["cohort"]=="strict" and r["channel"]==channel and r["method"]=="logit"}
        for cid in model_records:
            a, b = model_records[cid], baseline_records[cid]
            y = a["true_exit"]-1
            pa, pb = np.array(a["probabilities"]), np.array(b["probabilities"])
            paired.append({"case_id":cid, "date":a["test_date"], "channel":channel,
                           "llm_minus_logit_nll":float(-np.log(max(pa[y],1e-6))+np.log(max(pb[y],1e-6))),
                           "llm_correct":bool(pa.argmax()==y), "logit_correct":bool(pb.argmax()==y),
                           "llm_fallback":a["fallback"]})
    write_csv("paired_llm_vs_logit.csv", paired)
    lines = ["# 三出口真人选择数据清洗与首轮跨日期比较", "", "2026年10月12日。研究性离线评价，输入为记录的出口属性。", "",
             "## 完成的工作", "",
             f"核对14个三出口表头区块和PPT第3、5、10、15页。提取{summary['raw_rows']}行、{summary['person_keys']}个日期内人员编号。",
             f"主分析保留{len(strict)}个单条、单一数字选择记录；另保存{len(seq)}个有选择变化的最早记录，仅作时点未核实的探索分析。",
             "修正旧提取器对未写Level的标题的误归类。17b的描述明确为17的朋友，保留为不同编号，同时承认朋友及群体选择可能相关。", "",
             "| 日期 | 主分析记录 | 改道探索记录 |", "|---|---:|---:|"]
    if runtime_summary["complete"]:
        base = {r["method"]:r for r in aggregate_metrics if r["cohort"]=="strict" and r["channel"]=="primary" and r["test_date"]=="pooled"}
        lines[4:4] = [f"**本轮零样本Qwen 8B未超过真人校准的Logit。** 基础信息下，Qwen准确率{base['qwen3_8b']['accuracy']:.2%}、对数损失{base['qwen3_8b']['log_loss']:.4f}；Logit准确率{base['logit']['accuracy']:.2%}、对数损失{base['logit']['log_loss']:.4f}。这支持优先用Logit作为当前行为基线，尚不支持直接把未经校准的Qwen用作人群行为引擎。", ""]
    for d in DATES:
        lines.append(f"| {d} | {summary['strict_per_date'].get(d,0)} | {summary['sequence_per_date'].get(d,0)} |")
    lines += ["", "## 清洗与字段边界", "",
              "主分析剔除工作人员、空白及编码不明记录，并将同一人前后记录放到探索序列。排除记录均保留原行号及原因，没有补造缺失值。",
              "Staff=2出现在工作人员记录中，未自动改成普通疏散者的Guide=1。其他原始0/1编码仅在附加条件分析中使用。",
              "Fam在此楼层恒为[1,0,0]，与E1出口截距完全重合，不能单独估计熟悉程度的影响。Logit用出口截距吸收该固定差异。",
              "距离为近似欧氏距离，人群变量为此前选择该出口的人数。记录时点未逐人用视频验证，也没有完整个人感知信息。",
              "Excel与PPT对应差异：06-03-19的68、71为Excel空白编号，PPT无对应编号；31-10-19的PPT有76号，但Excel对应区块无该编号。其他日期编号相符。没有补造缺失数据。", "",
              "原论文表3对GF四日期报告19、37、21、44条选择，总计121条。本轮118个主/探索候选与其差3个带Staff人物描述的编号；这是本轮保守筛选，尚未证明与原作者排除管理工作人员的规则逐条一致。",
              f"主分析实际选择计数为E1={sum(r['initial_label']==1 for r in strict)}、E2={sum(r['initial_label']==2 for r in strict)}、E3={sum(r['initial_label']==3 for r in strict)}。没有出口3选择者，无法评价对此类人的识别能力。",
              "单条数字记录只说明表中没有该人的改道序列，首次决策时点仍未用录像逐人核实。", "",
              "## 冻结的比较方法", "",
              "每轮留出一个日期，其余三个日期训练；最近出口的概率温度和Logit正则强度在训练日期内再按日期留出选择。没有用外层日期调参。",
              "主比较使用距离、此前选择人数和固定Fam代理；附加比较加入原始二元Staff，并明确其时点未经逐人确认。",
              "最近出口准确率由最近距离决定（并列取最小编号），概率由训练内校准的距离softmax给出。出口频率采用训练集加一平滑。",
              "Logit含E1/E2截距、距离和此前选择人数，附加模型含Guide。所有拟合均收敛，系数仅用于预测，不报告因果解释。",
              "Qwen使用固定本地qwen3:8b，零样本、固定提示词和结构化输出。模型不接收真实选择、人物描述、日期类型或改道备注。",
              "主要指标为对数损失（越低越好），同时报告Brier概率误差、准确率和出口比例总变差。仅四个事件，不把人员数当作独立事件重复。", "",
              "准确率的概率并列按最小出口编号处理。均匀概率参照在此规则下始终取E1，其准确率没有随机猜测的含义，因此表内不报告该数值。四个日期等权均值另列，避免人数多的日期掩盖其他日期。", "",
              "## 主分析结果", ""]
    for channel, name in [("primary", "基础信息"), ("guide_conditional", "附加工作人员引导编码（条件性）")]:
        lines += [f"### {name}", "", "| 方法 | 样本 | 准确率 | 对数损失 | Brier | 出口比例偏差 | 回退 |",
                  "|---|---:|---:|---:|---:|---:|---:|"]
        for r in aggregate_metrics:
            if r["cohort"]=="strict" and r["channel"]==channel and r["test_date"]=="pooled":
                acc = "—" if r["method"]=="uniform" else f"{r['accuracy']:.2%}"
                lines.append(f"| {r['method']} | {r['n']} | {acc} | {r['log_loss']:.4f} | {r['brier']:.4f} | {r['exit_share_tv']:.2%} | {r['fallbacks']} |")
        lines += ["", "| 方法 | 四日期等权准确率 | 四日期等权对数损失 |", "|---|---:|---:|"]
        for method in ["training_frequency", "nearest_exit", "logit", "qwen3_8b"]:
            rr = [r for r in aggregate_metrics if r["cohort"]=="strict" and r["channel"]==channel and r["test_date"]!="pooled" and r["method"]==method]
            if len(rr)==4:
                lines.append(f"| {method} | {np.mean([r['accuracy'] for r in rr]):.2%} | {np.mean([r['log_loss'] for r in rr]):.4f} |")
        lines += ["", "| 留出日期 | 方法 | 准确率 | 对数损失 | 出口比例偏差 |", "|---|---|---:|---:|---:|"]
        for r in sorted(aggregate_metrics, key=lambda r:(r["test_date"],r["method"])):
            if r["cohort"]=="strict" and r["channel"]==channel and r["test_date"]!="pooled" and r["method"] in ("logit", "qwen3_8b", "training_frequency"):
                lines.append(f"| {r['test_date']} | {r['method']} | {r['accuracy']:.2%} | {r['log_loss']:.4f} | {r['exit_share_tv']:.2%} |")
        lines.append("")
    lines += ["## 最早改道记录的探索分析", "",
              "以下将15个选择变化记录的最早行加入103条主记录，初始标签取注释的第一个出口。原始属性的决策时点未核实，不能当作动态改道预测验证。", "",
              "| 方法 | 样本 | 准确率 | 对数损失 | Brier |", "|---|---:|---:|---:|---:|"]
    for r in aggregate_metrics:
        if r["cohort"]=="sequence_exploratory" and r["channel"]=="primary" and r["test_date"]=="pooled" and r["method"]!="uniform":
            lines.append(f"| {r['method']} | {r['n']} | {r['accuracy']:.2%} | {r['log_loss']:.4f} | {r['brier']:.4f} |")
    lines.append("")
    lines += ["## 执行状态与解释限制", "", f"LLM请求保存{len(llm_records)}/{expected_requests}条，有效{runtime_summary['valid_requests']}条，失败{runtime_summary['invalid_requests']}条。"]
    if runtime_summary["median_wall_seconds"] is not None:
        lines.append(f"单条请求中位耗时{runtime_summary['median_wall_seconds']:.2f}秒，请求耗时合计{runtime_summary['total_wall_seconds_sum']/60:.2f}分钟。221个请求包含同一案例的不同输入条件，不能称为221名独立受试者。")
    if not runtime_summary["complete"]:
        lines.append("LLM尚未完成全队列，表中的局部LLM结果不能与全样本基线直接比较。")
    else:
        for channel in ("primary", "guide_conditional"):
            rows = {r["method"]: r for r in aggregate_metrics if r["cohort"]=="strict" and r["channel"]==channel and r["test_date"]=="pooled"}
            a,b = rows["qwen3_8b"],rows["logit"]
            lines.append(f"{channel}：LLM相对Logit准确率差{(a['accuracy']-b['accuracy'])*100:+.2f}个百分点，对数损失差{a['log_loss']-b['log_loss']:+.4f}。正的损失差表示LLM概率预测更差。")
        first = next(r for r in aggregate_metrics if r["cohort"]=="strict" and r["channel"]=="primary" and r["method"]=="qwen3_8b" and r["test_date"]=="pooled")
        guided = next(r for r in aggregate_metrics if r["cohort"]=="strict" and r["channel"]=="guide_conditional" and r["method"]=="qwen3_8b" and r["test_date"]=="pooled")
        lines += ["", "## 保存回答揭示的问题", "",
                  f"真人记录中E1占78.64%、E2占21.36%、E3为0。本轮基础Qwen平均给E1分配{first['predicted_shares'][0]:.2%}、E2分配{first['predicted_shares'][1]:.2%}、E3分配{first['predicted_shares'][2]:.2%}的概率，预测分流比例与样本相差较大。不能因此说E3不合法或人在其他情况下不会选E3。",
                  "基础Qwen在81个实际选E1的记录中预测对51个，19个误判为E2、11个误判为E3；在22个实际选E2的记录中预测对12个。原始回答可用于核对这些错误，无需重新请求模型。",
                  f"加入Staff编码后，Qwen准确率提高{(guided['accuracy']-first['accuracy'])*100:.2f}个百分点，但实际选E2的正确识别从12个降到10个；整体改善主要来自更多正确预测E1，不能解释为已经学会服从工作人员或动态改道。",
                  "四个日期中Qwen的对数损失均高于相同信息条件的Logit。这是本轮设置下的描述性结果，不是所有LLM、所有提示或建筑的结论。", "",
                  "## 下一步的有限工作", "",
                  "优先核对15个序列中首次位置、人群计数和工作人员提醒的先后，并核对3个Staff描述编号是否属于原作者排除的管理人员。当前不扩大PPO训练。",
                  "如果后续检验真人示例或概率校准对Qwen的帮助，应另写冻结方案，在训练日期内完成适应，并明确本批日期已经用于探索，不能继续当作未接触的最终确认数据。",
                  "Logit目前可作为离线行为基线；进入动态仿真前仍需检查行为时点、运动模型和引导响应假设。", ""]
    lines += ["", "主分析只涵盖单条、无改道注释的记录，不能推广到所有人或动态改道。附加改道分析的最早属性与首次选择对应仍需视频或代码本确认。",
              "这一轮评估零样本Qwen与用其他日期真人数据拟合的Logit，代表各自实际部署方式；没有检验真人示例或校准是否能改善LLM。",
              "当前数据已被初审，属于探索评价。没有新增PPO训练、运动/烟气模拟，也没有读取48个封存最终疏散场景。",
              "原附件和既有模型选择、开发状态文件的哈希均再次核对。原研究为演习与误报警，结果不直接证明真实火灾控制效果。", "",
              "## 文件与来源", "",
              "清洗表：clean_choice_long.csv；逐行审计：row_audit.csv；排除原因：exclusion_reasons.csv；冻结方案：frozen_plan.json。",
              "所有方法逐人输出：all_predictions.csv；分日期指标：all_metrics.csv；训练内调参：baseline_fit_details.json；LLM完整请求响应：responses/。",
              "[原论文：Tong等，Safety Science 182 (2025), 106708](https://doi.org/10.1016/j.ssci.2024.106708)。方法3.3和图1提供字段及楼层对应。", ""]
    report_path = ROOT / "reports/Exit_choice三出口数据清洗与首轮实验_2026-10-12.md"
    report_path.write_text("\n".join(lines), encoding="utf-8")
    check_preservation()
    print(json.dumps({"report":str(report_path), **runtime_summary}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["prepare", "baselines", "llm", "report"])
    args = parser.parse_args()
    globals()[args.command]()

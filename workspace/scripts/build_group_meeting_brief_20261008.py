"""Create a compact group-meeting report from completed research records."""
from pathlib import Path
import hashlib
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf as fitz

from build_report_docx import configure_styles, set_run_font, add_page_number
from build_full_history_report import inline, table, split_table_row
from build_integrated_report_documents_20261008 import export_pdf

ROOT=Path(__file__).resolve().parents[1]
NAME='LLM_PPO火灾疏散研究_组会简报_2026-10-08'
ASSETS=ROOT/'artifacts/group_meeting_brief_assets_20261008'
CONTENT='''# LLM＋PPO火灾疏散研究：组会简报

**阶段性实验汇报｜截至2026年10月8日｜建议汇报约5分钟**

**当前结论：已完成仿真平台、对照实验和失败机制诊断；尚未证明LLM建议能稳定改善疏散效果。**

## 1. 研究问题与系统分工

研究问题是：在火灾、拥堵、出口封闭或预测错误时，加入LLM建议是否优于纯规则或PPO？系统能否判断当前建议是否值得采纳？

场景是两个房间、一个公共路口、两条出口走廊的五分区模型，默认240人，每步代表5秒。近期五个参数场景共用这一布局，烟气危险与人群参数主要为合成设定。[S1]

PPO即近端策略优化，可以理解为通过反复仿真、根据结果反馈学习控制的强化学习算法。本项目中，它根据人数、危险和出口状态，选择出口指示与房间放行档位；后期共9种请求动作：北/南指示各配25%、50%、75%、100%放行，以及暂停房间放行。它不直接决定每个人的路线，暂停也不停止已在走廊中的人员。

LLM提出高层候选；门控检查时效、执行条件与采纳依据；共同约束检查最终动作。近期执行逻辑为：现场观测→PPO基础动作与LLM候选→约束及门控→环境执行→新观测。人员选择与危险演化由独立仿真模块实现。

## 2. 已实际完成的七类实验

| 实验 | 做了什么 | 主要发现 |
|---|---|---|
| 1. PPO基础训练 | 训练动态指示与联合放行控制，用多个训练种子比较 | 策略能运行和学习；单次好成绩不足以证明稳定效果 |
| 2. 行为与噪声测试 | 改变指定人群行为模型及观测噪声，交叉测试 | 策略对人群设定敏感；此处是合成条件压力测试 |
| 3. 公开真人选择验证 | Qwen3 4B与频率、规则、Logit比较；三轮存档214条响应 | Snopková测试准确率61.65%，同任务频率基线86.59%；真人行为有效性尚未建立 |
| 4. LLM＋PPO公平对照 | 对应轮次统一动作、基础信息和训练预算 | 首轮LLM失能率比纯PPO高1.34个百分点；护栏轮未建立稳定改善 |
| 5. 动态冲突与联合控制 | 测试拥堵、短/长期封闭、预测错误，联合控制方向与放行 | 较低烟气出口也可能因低吞吐、滞留而增加暴露 |
| 6. 反事实教师与收益预测 | 同状态模拟候选计划，训练监督学生与选择器，审计候选库 | 存在离线规划机会；学生未稳定保留收益，教师成绩不等于在线LLM效果 |
| 7. 再次覆盖PPO门控 | 比较持续建议、保留放行档位、首次覆盖和公开状态门控 | 有害拒绝与有益误拒并存，当前门控未可靠识别收益 |

来源：S1逐轮记录。各轮条件不同，不能将成绩直接拼接成持续提升曲线。214条仅指三轮行为验证响应，不是全项目LLM请求总数。

## 3. 最近一轮完整实验：做到了什么

近期控制使用Qwen3 14B既有缓存、两种人群行为、五套既有PPO权重、即时与实测延迟建议。五个同图参数场景各使用四批新人口；本轮新增PPO训练和LLM请求均为0。[S2]

| 完成项 | 数量 | 怎样理解 |
|---|---:|---|
| 完整疏散回合 | 2,080 | 从初始状态运行至回合终点 |
| 反事实终点模拟 | 2,480 | 从保存状态强制不同动作，再运行到终点 |
| 建议机会 / 实际控制步 | 248 / 102,632 | 记录是否覆盖及最终执行动作 |

“覆盖”指用候选动作替换同状态PPO基础动作。此处无建议PPO是同权重消融，并非重新训练的纯PPO。反事实结果用于离线评价，不作为在线输入；标签按冻结标准与有限配对未来计算。

**即时模式24个再次覆盖机会：13个有害、5个有益、6个收益微小或近于中性。两种门控拒绝全部24个，因此既拦住13/13有害，也误拒5/5有益。**[S2]

这个结果尚未说明门控会判断收益。实测延迟模式没有再次覆盖机会，不能用它验证再次识别能力。局部反事实收益也不能直接累加成完整回合效果。

## 4. 当前判断与主要限制

第一轮公平对照失能差为+1.34个百分点，区间[+0.67，+2.10]；第二轮护栏方案为+1.02个百分点，区间[-1.57，+4.55]。正差表示失能更多；第二轮区间跨零，未证明稳定改善，也未证明等效。[S1，第7—8章]

目前存在方向、容量与采纳时机的交互，也发现公开观测与已记录控制历史相同而收益相反的病例。[S2] 简单增加模型规模或训练量，尚无证据能解决这些问题。

大量人口、权重和随机未来不是大量独立建筑。当前结果限定于合成环境，尚未完成真实建筑、物理烟气耦合和连续行人轨迹验证；约束违规为0也不代表人员失能为0。

## 5. 下一步与组会讨论

下一步先核查新增响应观测是否提供额外信息；比较“当前观测、普通历史、响应特征”三种收益预测，排除未来信息泄漏，冻结设置后用新人口验证。评价同时报告有害拦截、有益误拒、覆盖次数，以及完整回合的失能、危险暴露、疏散与拥堵。

人群响应收益预测目前只有代码初稿，尚未测试、拟合或验证。[S3] 继续暂缓扩大PPO训练，48个保留最终场景继续封存。

建议组会讨论：是否将主问题收敛为“不确定人群响应下，何时采纳高层建议”？最少需要增加哪些现实可测的观测，才能判断具体候选收益？

## 证据入口（供会后查阅）

[S1：全程整合报告](C:/Users/摩天轮大盗/Documents/ChatGPT/LLM探索/reports/LLM_RL火灾疏散研究全程整合报告_2026-10-08.md)；[S2：最近门控实验与独立审计](C:/Users/摩天轮大盗/Documents/ChatGPT/LLM探索/reports/再次覆盖PPO_结果解释与独立审计_2026-10-07.md)；[S3：尚未验证的预测器初稿](C:/Users/摩天轮大盗/Documents/ChatGPT/LLM探索/src/fire_evac/crowd_response_gain.py)。本次只整理已有材料，没有新增实验。
'''


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def plain(s):return re.sub(r'\*\*','',re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',s))
def norm(s):return re.sub(r'\s+','',s)


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    sources=[ROOT/'reports/LLM_RL火灾疏散研究全程整合报告_2026-10-08.md',ROOT/'reports/再次覆盖PPO_结果解释与独立审计_2026-10-07.md',ROOT/'src/fire_evac/crowd_response_gain.py']
    protected=sources+[ROOT/'outputs/local_llm_selection.json',ROOT/'outputs/third_round_development_status.json']
    for base in ('artifacts/LLM_RL火灾疏散研究全程整合报告_2026-10-08','artifacts/LLM_RL火灾疏散研究_5分钟简洁版_2026-10-08'):
        protected.extend(ROOT/(base+ext) for ext in ('.docx','.pdf'))
    before={str(p):sha(p) for p in protected}
    ASSETS.mkdir(parents=True,exist_ok=True)
    md=ROOT/'reports'/(NAME+'.md');docx=ROOT/'artifacts'/(NAME+'.docx');pdf=docx.with_suffix('.pdf')
    assert not any(p.exists() for p in (md,docx,pdf))
    md.write_text(CONTENT,encoding='utf-8')
    doc=Document();configure_styles(doc);s=doc.sections[0]
    s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.62);s.left_margin=s.right_margin=Inches(.67)
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.first_line_indent=Pt(21)
    heading=doc.styles['Heading 1'];heading.font.size=Pt(12)
    heading.paragraph_format.space_before=Pt(9);heading.paragraph_format.space_after=Pt(5)
    lines=CONTENT.splitlines();expected=[];i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_after=Pt(7)
            set_run_font(p.add_run(line[2:]),chinese='微软雅黑',size=18,bold=True)
            expected.append(line[2:]);continue
        if line.startswith('## '):
            p=doc.add_heading(line[3:],1)
            if line.startswith('## 3.'):p.paragraph_format.page_break_before=True
            if line=='## 证据入口（供会后查阅）':
                for r in p.runs:set_run_font(r,chinese='微软雅黑',size=9,bold=True)
            expected.append(line[3:]);continue
        if line.startswith('|'):
            rows=[]
            while True:
                cells=split_table_row(line)
                if not all(re.fullmatch(r':?-{2,}:?',c) for c in cells):
                    rows.append(cells);expected.extend(plain(c) for c in cells)
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                line=lines[i].strip();i+=1
            table(doc,rows,6.93);continue
        p=doc.add_paragraph()
        if line.startswith(('**阶段','来源：','[S1：')):
            p.paragraph_format.first_line_indent=Pt(0);inline(p,line,8.5)
        else:inline(p,line,10.5)
        expected.append(plain(line))
    header=s.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run('LLM＋PPO火灾疏散｜组会简报｜2026-10-08'),size=8,color='666666')
    add_page_number(s.footer.paragraphs[0])
    doc.core_properties.title='LLM＋PPO火灾疏散研究：组会简报'
    doc.core_properties.subject='七类已完成实验、近期结果、限制与待验证工作'
    doc.save(docx);export_pdf(docx)
    with zipfile.ZipFile(docx) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        word=''.join(n.text or '' for n in xml.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    d=fitz.open(pdf);body='\n'.join(p.get_text() for p in d)
    assert all(norm(v) in norm(word) and norm(v) in norm(body) for v in expected)
    assert '\ufffd' not in body and len(d)<=3
    for i,page in enumerate(d):
        assert page.get_text().strip()
        for b in page.get_text('blocks'):
            assert b[0]>=-1 and b[1]>=-1 and b[2]<=page.rect.width+1 and b[3]<=page.rect.height+1
        page.get_pixmap(matrix=fitz.Matrix(1.3,1.3),alpha=False).save(ASSETS/f'page_{i+1:02d}.png')
    assert all(sha(Path(p))==h for p,h in before.items())
    main_text=CONTENT.split('## 证据入口')[0]
    han=len(re.findall(r'[\u4e00-\u9fff]',main_text))
    qa=dict(pages=len(d),main_chinese_characters=han,text_units_checked=len(expected),table_count=len(doc.tables),preserved_files=before,missing_text=0,empty_pages=0,outside_page_text=0,new_experiments=0,new_ppo_training=0,new_llm_requests=0,reserved_final_scenarios_read=0,source_files=[str(p) for p in sources],artifacts=[dict(path=str(p),bytes=p.stat().st_size,sha256=sha(p)) for p in (md,docx,pdf)])
    (ASSETS/'qa_manifest.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in qa.items() if k not in ('preserved_files','artifacts','source_files')},ensure_ascii=False),flush=True)


if __name__=='__main__':main()
